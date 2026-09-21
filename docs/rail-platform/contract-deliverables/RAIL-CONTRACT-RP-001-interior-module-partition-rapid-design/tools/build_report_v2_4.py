#!/usr/bin/env python3
"""Build a new V2.4 review revision without altering prior deliverables."""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

import build_report_v2_3 as layout

BASE = Path(__file__).resolve().parents[1]
SOURCE = BASE / 'source' / 'REPORT_V2.4.md'
OUTPUT = BASE / 'deliverables' / '基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告_V2.4.docx'


def cover(doc):
    for _ in range(4):
        p = doc.add_paragraph(); layout.set_fixed_24(p, first_line=False)
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(36)
    p.paragraph_format.space_after = Pt(24)
    layout.format_style(doc.styles['Title'], layout.H1_LATIN, layout.H1_CJK, 24, bold=True)
    for element in (doc.styles['Title'].element, p._p):
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    layout.format_run(p.add_run('基于轨道交通客运装备的\n内装模块分区识别与快速设计\n系统构建研究报告'), layout.H1_LATIN, layout.H1_CJK, 24, bold=True)
    for value in ['平台构建技术指导', '', '文档版本  V2.4', '委托单位  中车工业研究院有限公司', '承担单位  北京科技大学', '', '2026年9月18日']:
        p = doc.add_paragraph(); layout.set_fixed_24(p, first_line=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        layout.format_run(p.add_run(value), size=12)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    layout.configure_document(doc); layout.configure_header_footer(section)
    for p in doc.sections[0].footer.paragraphs: p.clear()


def build():
    parser = argparse.ArgumentParser()
    parser.add_argument('--page-map', type=Path, default=SOURCE.with_name('TOC_V2.4.json'))
    parser.add_argument('--out', type=Path, default=OUTPUT)
    args = parser.parse_args()
    layout.SOURCE_PATH = SOURCE
    layout.TOC_PAGE_MAP = json.loads(args.page_map.read_text()) if args.page_map.exists() else {}
    source = SOURCE.read_text(encoding='utf-8')
    entries = layout.extract_toc_entries(source)
    doc = Document(); layout.configure_document(doc)
    doc.core_properties.title = layout.DOC_TITLE
    doc.core_properties.subject = '内装模块划分、识别方法、快速设计流程及平台构建方案'
    doc.core_properties.author = '轨道客室智能设计平台项目组'
    doc.core_properties.comments = 'RAIL-CONTRACT-RP-001 V2.4'
    cover(doc); layout.add_front_matter(doc, entries)
    layout.parse_markdown(doc, source, entries)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == 'Heading 1' and paragraph.text == '参考文献':
            paragraph.paragraph_format.page_break_before = True
    # Use the native Word comments API so local revision explanations stay at
    # their relevant passages; original reviewer comments are separately archived.
    notes = [
        ('出口客运装备内装设计需要在产品族共性', 'R01：摘要改为概述，删除关键词；概括背景、方法、路线及平台构建成果。'),
        ('项目面向中车出口产品平台开展研究', 'R02：删除原第一章，必要研究对象说明移入本节，合同节点保留在计划与验证章节。'),
        ('项目技术方案将不同国家', 'R03：补充出口产品平台、设计流程问题和 AI 技术分析。现状采用已有技术资料分析，企业流程、访谈及统计记录仍需甲方补充确认。'),
        ('平台架构承接前文', 'R04：平台构建移至方法与快速设计流程之后；新增五层逻辑总图，按层说明建设思路。用户层是在技术方案原分层上显式展开的使用对象。'),
        ('传统模块化划分首先确定', 'R05：补充功能、结构、接口、维护和产品族划分，明确组合方法。具体车型 BOM 与模块边界需专业资料确认。'),
        ('检测方法以目标框表达', 'R06：补充识别原理与比较，明确当前标记驱动编辑路线及下一阶段提示分割候选路线，未将候选方法写成已部署算法。'),
        ('颜色分区设计从项目原图开始', 'R07：增加实际工作流及条件图生成过程。尚缺同一真实客室案例的原图、标记图、指令、生成结果和评价；本版不以流程图冒充效果验证，本条部分响应。'),
        ('本研究围绕出口客运装备内装的共性模块复用', 'R08：结论按现状与模块化、方法选型、技术流程及平台架构组织，只总结已形成的分析与方案。'),
        ('平台构建方案按用户、应用、软件支撑', 'R09：正式稿删除两个附录，原文保留于历史版本；合同追踪与待补材料转内部响应清单，结论后保留技术参考文献。'),
    ]
    for anchor, text in notes:
        matches = [p for p in doc.paragraphs if p.text.startswith(anchor)]
        if len(matches) != 1: raise ValueError(f'Comment anchor not unique: {anchor}')
        doc.add_comment(matches[0].runs, text=text, author='修订说明', initials='RP')
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(args.out)


if __name__ == '__main__':
    build()
