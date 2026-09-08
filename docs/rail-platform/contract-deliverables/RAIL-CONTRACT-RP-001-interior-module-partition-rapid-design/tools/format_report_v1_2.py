#!/usr/bin/env python3
"""Build the V1.2 format-only edition from the V1.1 report content."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_report_docx as base


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_PATH = (
    BASE_DIR
    / "deliverables"
    / "基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告_V1.2.docx"
)

DOC_VERSION = "V1.2"
FINAL_PAGE_COUNT = 60

A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7
TOP_MARGIN_CM = 2.5
BOTTOM_MARGIN_CM = 2.5
LEFT_MARGIN_CM = 2.6
RIGHT_MARGIN_CM = 2.4
HEADER_DISTANCE_CM = 1.5
FOOTER_DISTANCE_CM = 1.75
CONTENT_WIDTH_DXA = 9072
TABLE_INDENT_DXA = 100
TABLE_WIDTH_DXA = CONTENT_WIDTH_DXA

BODY_FONT = "SimSun"
BODY_FONT_EAST_ASIA = "宋体"
HEADING_1_FONT = "SimHei"
HEADING_1_FONT_EAST_ASIA = "黑体"
HEADING_2_FONT = "KaiTi"
HEADING_2_FONT_EAST_ASIA = "楷体"
BLACK = RGBColor(0, 0, 0)


CHINESE_NUMERALS = {
    1: "一",
    2: "二",
    3: "三",
    4: "四",
    5: "五",
    6: "六",
    7: "七",
    8: "八",
    9: "九",
    10: "十",
    11: "十一",
    12: "十二",
    13: "十三",
    14: "十四",
    15: "十五",
    16: "十六",
    17: "十七",
    18: "十八",
    19: "十九",
    20: "二十",
}


def set_rfonts(target, latin: str, east_asia: str):
    rpr = target._element.get_or_add_rPr() if hasattr(target, "_element") else target
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def format_run(run, latin: str, east_asia: str, size: float, *, bold=None):
    run.font.name = latin
    set_rfonts(run, latin, east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def format_style(style, latin: str, east_asia: str, size: float, *, bold=None):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    if bold is not None:
        style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)


def set_fixed_24(paragraph, *, first_line: bool):
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(24 if first_line else 0)


def has_numbering(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def has_inline_code(paragraph) -> bool:
    return any(
        run._r.rPr is not None and run._r.rPr.find(qn("w:shd")) is not None
        for run in paragraph.runs
    )


def replace_in_runs(paragraph, old: str, new: str):
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)


def chinese_level_one(text: str) -> str:
    match = re.match(r"^(\d+)\s+(.+)$", text)
    if not match:
        return text
    number = int(match.group(1))
    numeral = CHINESE_NUMERALS.get(number, str(number))
    return f"{numeral}、{match.group(2)}"


def chinese_level_two(text: str) -> str:
    match = re.match(r"^\d+\.(\d+)\s+(.+)$", text)
    if not match:
        return text
    number = int(match.group(1))
    numeral = CHINESE_NUMERALS.get(number, str(number))
    return f"（{numeral}）{match.group(2)}"


def scale_table_widths(table):
    grid = table._tbl.tblGrid
    widths = [int(col.get(qn("w:w"))) for col in grid.gridCol_lst]
    if not widths or sum(widths) <= 0:
        widths = [1] * len(table.columns)
    scaled = [max(240, round(width / sum(widths) * TABLE_WIDTH_DXA)) for width in widths]
    scaled[-1] += TABLE_WIDTH_DXA - sum(scaled)
    base.apply_table_geometry(
        table,
        scaled,
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa={"top": 60, "bottom": 60, "start": 100, "end": 100},
    )


def update_version_metadata(doc: Document):
    control = doc.tables[0]
    control.cell(1, 1).text = f"{base.DOC_ID} / {DOC_VERSION}"

    revisions = doc.tables[1]
    revisions.cell(2, 0).text = "V1.1"
    revisions.cell(2, 1).text = "2026-09-03"
    revisions.cell(2, 2).text = base.CODE_SHORT
    revisions.cell(2, 3).text = "历史评审稿"
    revisions.cell(2, 4).text = "依据技术开发合同（26版）扩充研究内容、方法、实验、指标、里程碑和追踪矩阵"
    values = [
        DOC_VERSION,
        "2026-09-03",
        base.CODE_SHORT,
        "当前格式评审稿",
        "正文内容不变；按指定 A4、页边距、字体字号、固定 24 磅行距、首行缩进和标题编号格式调整",
    ]
    row = revisions.add_row()
    for cell, value in zip(row.cells, values):
        cell.text = value
    base.set_row_cant_split(row)

    for paragraph in doc.paragraphs:
        replace_in_runs(paragraph, "文档版本：V1.1", f"文档版本：{DOC_VERSION}")
        replace_in_runs(paragraph, "V1.1 编制时", "V1.2 沿用 V1.1 编制时")
        replace_in_runs(paragraph, "当前 V1.1 已完成", "当前 V1.2 沿用 V1.1 已完成的")
        replace_in_runs(paragraph, "当前 V1.1 的局限", "当前 V1.2 沿用 V1.1 的研究局限")
        replace_in_runs(paragraph, "不属于 V1.1 报告事实范围", "不属于 V1.2 报告事实范围")
        replace_in_runs(
            paragraph,
            "V1.1 相对 V1.0 的变化仅为依据技术开发合同（26版）完善研究内容、技术路线、指标口径和交付计划。",
            "V1.2 相对 V1.1 仅调整文章格式，研究内容、事实基线与结论不变。",
        )

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                if "56 页逐页复核" in text:
                    page_text = (
                        f"V1.2 共 {FINAL_PAGE_COUNT} 页并完成逐页复核；"
                        if FINAL_PAGE_COUNT
                        else "V1.2 格式版交付前执行逐页复核；"
                    )
                    cell.text = (
                        page_text
                        + "页面参数、表格几何、标题、图片及无障碍检查通过，样式提示逐项复核"
                    )
                elif "V1.1 源稿、DOCX、图表与版本记录" in text:
                    cell.text = "V1.1 内容源稿、V1.2 格式说明与 DOCX、图表及版本记录"


def configure_page_and_styles(doc: Document):
    for section in doc.sections:
        section.page_width = Cm(A4_WIDTH_CM)
        section.page_height = Cm(A4_HEIGHT_CM)
        section.top_margin = Cm(TOP_MARGIN_CM)
        section.bottom_margin = Cm(BOTTOM_MARGIN_CM)
        section.left_margin = Cm(LEFT_MARGIN_CM)
        section.right_margin = Cm(RIGHT_MARGIN_CM)
        section.gutter = Cm(0)
        section.header_distance = Cm(HEADER_DISTANCE_CM)
        section.footer_distance = Cm(FOOTER_DISTANCE_CM)

    normal = doc.styles["Normal"]
    format_style(normal, BODY_FONT, BODY_FONT_EAST_ASIA, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(24)
    normal.paragraph_format.first_line_indent = Pt(24)

    list_style = doc.styles["List Paragraph"]
    format_style(list_style, BODY_FONT, BODY_FONT_EAST_ASIA, 12)
    list_style.paragraph_format.space_before = Pt(0)
    list_style.paragraph_format.space_after = Pt(0)
    list_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    list_style.paragraph_format.line_spacing = Pt(24)
    list_style.paragraph_format.first_line_indent = Pt(0)

    heading_styles = {
        "Heading 1": (HEADING_1_FONT, HEADING_1_FONT_EAST_ASIA, True),
        "Heading 2": (HEADING_2_FONT, HEADING_2_FONT_EAST_ASIA, False),
        "Heading 3": (BODY_FONT, BODY_FONT_EAST_ASIA, False),
        "Heading 4": (BODY_FONT, BODY_FONT_EAST_ASIA, False),
    }
    for name, (latin, east_asia, bold) in heading_styles.items():
        style = doc.styles[name]
        format_style(style, latin, east_asia, 14, bold=bold)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(24)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    format_style(caption, BODY_FONT, BODY_FONT_EAST_ASIA, 12)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.line_spacing = Pt(24)
    caption.paragraph_format.first_line_indent = Pt(0)


def format_document_paragraphs(doc: Document):
    in_front_matter = False
    in_main_body = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""

        if style_name == "Heading 1" and text == "文档控制":
            in_front_matter = True
        if style_name == "Heading 1" and text == "摘要":
            in_main_body = True

        if text == "目录":
            paragraph.paragraph_format.page_break_before = True

        if in_front_matter and not in_main_body and re.match(r"^\d+\s+", text):
            paragraph.text = chinese_level_one(text)
            text = paragraph.text.strip()
        elif in_main_body and style_name == "Heading 1":
            paragraph.text = chinese_level_one(text)
            text = paragraph.text.strip()
        elif in_main_body and style_name == "Heading 2":
            paragraph.text = chinese_level_two(text)
            text = paragraph.text.strip()

        if not in_front_matter:
            continue

        if style_name.startswith("Heading"):
            level = int(style_name.rsplit(" ", 1)[-1])
            if level == 1:
                font = (HEADING_1_FONT, HEADING_1_FONT_EAST_ASIA, True)
            elif level == 2:
                font = (HEADING_2_FONT, HEADING_2_FONT_EAST_ASIA, False)
            else:
                font = (BODY_FONT, BODY_FONT_EAST_ASIA, False)
            for run in paragraph.runs:
                format_run(run, font[0], font[1], 14, bold=font[2])
            set_fixed_24(paragraph, first_line=False)
            continue

        if has_drawing(paragraph):
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.first_line_indent = Pt(0)
        elif has_numbering(paragraph):
            paragraph.style = doc.styles["List Paragraph"]
            set_fixed_24(paragraph, first_line=False)
        else:
            is_body_prose = in_main_body and bool(text) and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER
            set_fixed_24(paragraph, first_line=is_body_prose)
            if has_inline_code(paragraph):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            format_run(run, BODY_FONT, BODY_FONT_EAST_ASIA, 12)


def format_tables(doc: Document):
    for table in doc.tables:
        scale_table_widths(table)
        for row_index, row in enumerate(table.rows):
            base.set_row_cant_split(row)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    set_fixed_24(paragraph, first_line=False)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        format_run(
                            run,
                            BODY_FONT,
                            BODY_FONT_EAST_ASIA,
                            12,
                            bold=True if row_index == 0 else None,
                        )


def update_numbering_punctuation(doc: Document):
    numbering = doc.part.numbering_part.element
    for level in numbering.findall(f".//{qn('w:lvl')}"):
        num_fmt = level.find(qn("w:numFmt"))
        level_text = level.find(qn("w:lvlText"))
        if (
            num_fmt is not None
            and num_fmt.get(qn("w:val")) == "decimal"
            and level_text is not None
        ):
            level_text.set(qn("w:val"), "%1．")


def format_footer(doc: Document):
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            set_fixed_24(paragraph, first_line=False)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                format_run(run, BODY_FONT, BODY_FONT_EAST_ASIA, 9)


def format_cover(doc: Document):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() != base.DOC_TITLE:
            continue
        paragraph.text = "基于轨道交通客运装备的\n内装模块分区识别与快速设计\n系统构建研究报告"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            base.set_run_font(
                run,
                size=27,
                color=base.COVER_NAVY,
                bold=True,
                heading=True,
            )
        break


def build():
    base.DOC_VERSION = DOC_VERSION
    base.OUTPUT_PATH = OUTPUT_PATH
    base.build()

    doc = Document(OUTPUT_PATH)
    update_version_metadata(doc)
    configure_page_and_styles(doc)
    format_cover(doc)
    format_document_paragraphs(doc)
    format_tables(doc)
    update_numbering_punctuation(doc)
    format_footer(doc)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
