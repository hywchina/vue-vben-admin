#!/usr/bin/env python3
"""Build the formal V2.3 platform-construction guidance report."""

from __future__ import annotations

import re
import json
import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips




def column_widths_from_weights(weights, total):
    widths = [int(total*w/sum(weights)) for w in weights]
    widths[-1] += total-sum(widths)
    return widths


def apply_table_geometry(table, widths, *, table_width_dxa, indent_dxa, cell_margins_dxa):
    pr = table._tbl.tblPr
    tw = pr.find(qn("w:tblW"))
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(table_width_dxa))
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), str(indent_dxa)); ind.set(qn("w:type"), "dxa"); pr.append(ind)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for cell,width in zip(row.cells,widths):
            cell.width=Twips(width)
            tcpr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement("w:tcMar")
            for side,value in cell_margins_dxa.items():
                el=OxmlElement(f"w:{side}"); el.set(qn("w:w"),str(value)); el.set(qn("w:type"),"dxa"); mar.append(el)
            tcpr.append(mar)


BASE_DIR = Path(__file__).resolve().parents[1]
SOURCE_PATH = BASE_DIR / "source" / "REPORT_V2.3.md"
ASSETS_DIR = BASE_DIR / "assets"
OUTPUT_PATH = (
    BASE_DIR
    / "deliverables"
    / "基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告_V2.3.docx"
)

DOC_ID = "RAIL-CONTRACT-RP-001"
DOC_VERSION = "V2.3"
DOC_TITLE = "基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告"
SYSTEM_NAME = "客运装备内装模块化分区快速设计平台"
COMPILED_DATE = "2026年9月8日"

BODY_LATIN = "SimSun"
BODY_CJK = "宋体"
H1_LATIN = "SimHei"
H1_CJK = "黑体"
H2_LATIN = "KaiTi"
H2_CJK = "楷体"

BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 98, 108)
HEADER_FILL = "EEF1F5"
BORDER = "D9D9D9"
ACCENT = "B91C32"

CONTENT_WIDTH_DXA = 9072
TABLE_INDENT_DXA = 100
TABLE_WIDTH_DXA = CONTENT_WIDTH_DXA - TABLE_INDENT_DXA
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 100, "end": 100}

# Filled after the first deterministic render. The visible values make the
# table of contents useful even before Word refreshes fields; every entry also
# carries an internal hyperlink to its heading bookmark.
TOC_PAGE_MAP = {}


def set_rfonts(target, latin: str, east_asia: str):
    rpr = target._element.get_or_add_rPr() if hasattr(target, "_element") else target
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def format_run(run, latin: str = BODY_LATIN, east_asia: str = BODY_CJK, size: float = 12, *, bold=None):
    run.font.name = latin
    set_rfonts(run, latin, east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def format_style(style, latin: str, east_asia: str, size: float, *, bold=None):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    if bold is not None:
        style.font.bold = bold
    set_rfonts(style, latin, east_asia)


def set_fixed_24(paragraph, *, first_line: bool):
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(24 if first_line else 0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def set_cell_fill(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = BORDER, size: str = "4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:cantSplit")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def add_page_field(paragraph):
    paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, result, end):
        run = paragraph.add_run()
        run._r.append(item)
    paragraph.add_run(" 页")
    for run in paragraph.runs:
        format_run(run, size=9)


def add_numbering_definition(doc: Document, *, decimal: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if decimal else "bullet")
    level.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1．" if decimal else "•")
    level.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "480")
    ind.set(qn("w:hanging"), "240")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "480")
    spacing.set(qn("w:lineRule"), "exact")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:numPr"))
    if existing is not None:
        ppr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def add_inline(paragraph, text: str, *, bold=False):
    token_re = re.compile(r"(\*\*[^*]+\*\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            format_run(run, bold=bold)
        run = paragraph.add_run(match.group(0)[2:-2])
        format_run(run, bold=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        format_run(run, bold=bold)


def configure_document(doc: Document):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.4)
        section.gutter = Cm(0)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    normal = doc.styles["Normal"]
    format_style(normal, BODY_LATIN, BODY_CJK, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(24)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    list_style = doc.styles["List Paragraph"]
    format_style(list_style, BODY_LATIN, BODY_CJK, 12)
    list_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    list_style.paragraph_format.line_spacing = Pt(24)
    list_style.paragraph_format.first_line_indent = Pt(0)
    list_style.paragraph_format.space_before = Pt(0)
    list_style.paragraph_format.space_after = Pt(0)

    heading_styles = {
        "Heading 1": (H1_LATIN, H1_CJK, True),
        "Heading 2": (H2_LATIN, H2_CJK, False),
        "Heading 3": (BODY_LATIN, BODY_CJK, False),
        "Heading 4": (BODY_LATIN, BODY_CJK, False),
    }
    for name, (latin, cjk, bold) in heading_styles.items():
        style = doc.styles[name]
        format_style(style, latin, cjk, 14, bold=bold)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(24)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    format_style(caption, BODY_LATIN, BODY_CJK, 12, bold=False)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.line_spacing = Pt(24)
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)
    caption.paragraph_format.keep_with_next = True

    settings = doc.settings._element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def configure_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for paragraph in section.header.paragraphs:
        paragraph.clear()
    footer_p = section.footer.paragraphs[0]
    footer_p.clear()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_fixed_24(footer_p, first_line=False)
    add_page_field(footer_p)


def add_table(doc: Document, rows: list[list[str]], weights: list[float] | None = None):
    cols = max(len(row) for row in rows)
    normalized = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    if weights is None:
        lengths = [max(len(row[col]) for row in normalized) for col in range(cols)]
        weights = [max(5.0, min(float(length), 30.0)) ** 0.7 for length in lengths]
    widths = column_widths_from_weights(weights, TABLE_WIDTH_DXA)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        set_row_cant_split(row)
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                set_cell_fill(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 or len(normalized[row_idx][col_idx]) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            set_fixed_24(p, first_line=False)
            add_inline(p, normalized[row_idx][col_idx], bold=row_idx == 0)
    set_repeat_table_header(table.rows[0])
    return table


def extract_toc_entries(source: str) -> list[tuple[int, str, str]]:
    lines = source.splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1 :]
    entries = []
    for line in lines:
        match = re.match(r"^(#{1,3})\s+(.+)$", line.strip())
        if not match:
            continue
        level = len(match.group(1))
        text = match.group(2)
        entries.append((level, text, f"report_heading_{len(entries) + 1:03d}"))
    return entries


def add_internal_hyperlink(paragraph, text: str, anchor: str, *, bold=False):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = paragraph.add_run(text)
    format_run(run, size=12, bold=bold)
    run.font.color.rgb = BLACK
    run.font.underline = False
    paragraph._p.remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def add_heading_bookmark(paragraph, bookmark_name: str, bookmark_id: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(1, start)
    paragraph._p.append(end)


def add_clickable_toc(doc: Document, toc_entries: list[tuple[int, str, str]]):
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("目录")
    format_run(run, H1_LATIN, H1_CJK, 14, bold=True)

    # The printed directory is chapter-level; all subsection headings retain
    # outline levels and bookmarks for Word's navigation pane.
    for level, text, anchor in toc_entries:
        if level != 1:
            continue
        p = doc.add_paragraph()
        set_fixed_24(p, first_line=False)
        p.paragraph_format.left_indent = Pt(0 if level == 1 else 24 if level == 2 else 48)
        p.paragraph_format.right_indent = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.7),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )
        add_internal_hyperlink(p, text, anchor, bold=level == 1)
        tab = p.add_run("\t")
        format_run(tab)
        page = str(TOC_PAGE_MAP.get(anchor, "1"))
        add_internal_hyperlink(p, page, anchor, bold=level == 1)


def add_picture(doc: Document, image_path: Path, caption_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    # A drawing is taller than a text line. Applying the body style's exact
    # 24-point line box here clips the image in Word/LibreOffice.
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run()
    inline = run.add_picture(str(image_path), width=Inches(6.18))
    inline._inline.docPr.set("descr", caption_text)
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_with_next = False
    run = cap.add_run(caption_text)
    format_run(run)


def parse_markdown(doc: Document, source: str, toc_entries: list[tuple[int, str, str]]):
    lines = source.splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1 :]

    i = 0
    active_decimal = None
    heading_index = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for row_idx, table_line in enumerate(table_lines):
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if row_idx == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            active_decimal = None
            continue

        image_match = re.fullmatch(r"!\[(图\s*\d+\s+.+?)\]\((.+?)\)", stripped)
        if image_match:
            caption, rel_path = image_match.groups()
            add_picture(doc, (SOURCE_PATH.parent / rel_path).resolve(), caption)
            active_decimal = None
            i += 1
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            p = doc.add_heading(text, level=level)
            expected_level, expected_text, bookmark = toc_entries[heading_index]
            if expected_level != level or expected_text != text:
                raise ValueError(f"heading order mismatch: {text}")
            add_heading_bookmark(p, bookmark, heading_index + 1)
            heading_index += 1
            set_fixed_24(p, first_line=False)
            p.paragraph_format.space_before = Pt(8 if level == 1 else 4)
            latin, cjk, bold = {
                1: (H1_LATIN, H1_CJK, True),
                2: (H2_LATIN, H2_CJK, False),
                3: (BODY_LATIN, BODY_CJK, False),
                4: (BODY_LATIN, BODY_CJK, False),
            }[level]
            for run in p.runs:
                format_run(run, latin, cjk, 14, bold=bold)
            active_decimal = None
            i += 1
            continue

        if re.fullmatch(r"表\s*\d+\s+.+", stripped):
            p = doc.add_paragraph(style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped)
            format_run(run)
            active_decimal = None
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            if active_decimal is None:
                active_decimal = add_numbering_definition(doc, decimal=True)
            p = doc.add_paragraph(style="List Paragraph")
            set_fixed_24(p, first_line=False)
            apply_num(p, active_decimal)
            add_inline(p, numbered.group(1))
            i += 1
            continue

        if re.match(r"^-\s+", stripped):
            active_decimal = None
            bullet = add_numbering_definition(doc, decimal=False)
            p = doc.add_paragraph(style="List Paragraph")
            set_fixed_24(p, first_line=False)
            apply_num(p, bullet)
            add_inline(p, re.sub(r"^-\s+", "", stripped))
            i += 1
            continue

        active_decimal = None
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if (
                not candidate
                or candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith("![")
                or re.fullmatch(r"表\s*\d+\s+.+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
                or candidate.startswith("- ")
            ):
                break
            paragraph_lines.append(candidate)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        set_fixed_24(p, first_line=True)
        add_inline(p, " ".join(paragraph_lines))



def add_cover(doc):
    for _ in range(4):
        p=doc.add_paragraph(); set_fixed_24(p,first_line=False)
    p=doc.add_paragraph(style="Title")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing=Pt(36)
    p.paragraph_format.space_after=Pt(24)
    format_style(doc.styles["Title"],H1_LATIN,H1_CJK,24,bold=True)
    for element in (doc.styles["Title"].element, p._p):
        for border in list(element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    run=p.add_run("基于轨道交通客运装备的\n内装模块分区识别与快速设计\n系统构建研究报告")
    format_run(run,H1_LATIN,H1_CJK,24,bold=True)
    for value in ["平台构建技术指导", "", "文档版本  V2.3", "委托单位  中车工业研究院有限公司", "承担单位  北京科技大学", "", COMPILED_DATE]:
        p=doc.add_paragraph(); set_fixed_24(p,first_line=False)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(value),size=12)
    section=doc.add_section(WD_SECTION.NEW_PAGE)
    configure_document(doc); configure_header_footer(section)
    for p in doc.sections[0].footer.paragraphs: p.clear()


def add_front_matter(doc, entries):
    # Only substantive headings enter the clickable directory.
    add_clickable_toc(doc, entries)
    doc.add_page_break()


def build():
    global TOC_PAGE_MAP
    parser=argparse.ArgumentParser()
    parser.add_argument("--page-map",type=Path)
    args=parser.parse_args()
    page_map = args.page_map or SOURCE_PATH.with_name("TOC_V2.3.json")
    if page_map.exists(): TOC_PAGE_MAP=json.loads(page_map.read_text())
    source=SOURCE_PATH.read_text(encoding="utf-8")
    entries=extract_toc_entries(source)
    doc=Document(); configure_document(doc)
    doc.core_properties.title=DOC_TITLE
    doc.core_properties.subject="内装分区与快速设计平台的研究方法和构建指导"
    doc.core_properties.author="轨道客室智能设计平台项目组"
    doc.core_properties.keywords="轨道交通,内装模块,分区设计,快速设计,平台构建"
    doc.core_properties.comments="RAIL-CONTRACT-RP-001 V2.3"
    add_cover(doc); add_front_matter(doc,entries); parse_markdown(doc,source,entries)
    OUTPUT_PATH.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__": build()
