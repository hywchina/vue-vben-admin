#!/usr/bin/env python3
"""Build the formal V2.0 platform-construction guidance report."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

import build_report_docx as figure_base


SKILL_SCRIPTS = Path(
    "/home/huyanwei/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.826.12353/skills/documents/scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


BASE_DIR = Path(__file__).resolve().parents[1]
SOURCE_PATH = BASE_DIR / "source" / "REPORT_V2.0.md"
ASSETS_DIR = BASE_DIR / "assets"
OUTPUT_PATH = (
    BASE_DIR
    / "deliverables"
    / "基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告_V2.0.docx"
)

DOC_ID = "RAIL-CONTRACT-RP-001"
DOC_VERSION = "V2.0"
DOC_TITLE = "基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告"
SYSTEM_NAME = "客运装备内装模块化分区快速设计平台"
COMPILED_DATE = "2026年9月3日"

BODY_LATIN = "SimSun"
BODY_CJK = "宋体"
H1_LATIN = "SimHei"
H1_CJK = "黑体"
H2_LATIN = "KaiTi"
H2_CJK = "楷体"

BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 98, 108)
HEADER_FILL = "EEF1F5"
BORDER = "C9D1DA"
ACCENT = "B91C32"

CONTENT_WIDTH_DXA = 9072
TABLE_INDENT_DXA = 100
TABLE_WIDTH_DXA = CONTENT_WIDTH_DXA - TABLE_INDENT_DXA
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 100, "end": 100}

# Filled after the first deterministic render. The visible values make the
# table of contents useful even before Word refreshes fields; every entry also
# carries an internal hyperlink to its heading bookmark.
TOC_PAGE_MAP: dict[str, int] = {
    "摘要": 6,
    "一、报告定位与研究依据": 6,
    "（一）报告定位": 6,
    "（二）建设依据": 8,
    "（三）研究范围": 8,
    "（四）研究方法": 8,
    "二、建设背景与总体目标": 8,
    "（一）业务背景": 8,
    "（二）总体目标": 8,
    "（三）建设原则": 9,
    "三、总体研究框架与关键问题": 10,
    "（一）研究内容体系": 10,
    "（二）需要解决的关键问题": 11,
    "（三）研究成果形态": 11,
    "四、平台总体架构设计": 12,
    "（一）分层架构": 12,
    "（二）服务边界": 13,
    "（三）统一对象和业务协议": 13,
    "（四）部署与扩展架构": 13,
    "五、内装模块分区识别关键技术": 14,
    "（一）区域与模块对象模型": 14,
    "（二）多模态输入与预处理": 15,
    "（三）候选识别与多模态融合": 16,
    "（四）人机协同复核": 17,
    "（五）识别数据集与评价指标": 17,
    "六、多模态应用关键技术": 18,
    "（一）多模态能力定位": 18,
    "（二）语义对齐与知识增强": 18,
    "（三）能力编排": 18,
    "（四）模型选型与评测": 19,
    "七、快速设计系统关键技术": 19,
    "（一）模块化、参数化与规则化": 19,
    "（二）快速设计流程": 19,
    "（三）规则与约束体系": 20,
    "（四）跨区域协同优化": 21,
    "（五）方案检索、评价与复用": 21,
    "八、平台功能模块设计": 21,
    "（一）功能总体布局": 21,
    "（二）典型业务场景": 23,
    "（三）角色与权限": 23,
    "九、数据资源与知识底座": 23,
    "（一）数据资源体系": 23,
    "（二）数据治理流程": 25,
    "（三）数据质量控制": 25,
    "十、系统集成与接口设计": 25,
    "（一）与既有产品平台的交互": 25,
    "（二）外部智能能力接入": 26,
    "（三）接口验证": 26,
    "十一、安全、质量与运行保障": 26,
    "（一）安全与保密": 26,
    "（二）可追溯设计": 26,
    "（三）运行保障": 27,
    "十二、实施技术路线与进度计划": 27,
    "（一）总体实施路线": 27,
    "（二）阶段进度计划": 27,
    "（三）近期工作分解": 28,
    "（四）项目组织与沟通": 28,
    "十三、验证方法与验收建议": 29,
    "（一）分层验证方法": 29,
    "（二）典型验收场景": 30,
    "（三）质量指标处理原则": 30,
    "十四、主要风险与控制措施": 30,
    "（一）风险管理机制": 31,
    "十五、结论与建设建议": 31,
    "附录 A 甲方关注事项与报告章节对应表": 32,
    "附录 B 后续详细设计输入清单": 33,
}


def set_rfonts(target, latin: str, east_asia: str):
    rpr = target._element.get_or_add_rPr() if hasattr(target, "_element") else target
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def format_run(run, latin: str = BODY_LATIN, east_asia: str = BODY_CJK, size: float = 12, *, bold=None):
    run.font.name = latin
    set_rfonts(run, latin, east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def format_style(style, latin: str, east_asia: str, size: float, *, bold=None):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    if bold is not None:
        style.font.bold = bold
    set_rfonts(style, latin, east_asia)


def set_fixed_24(paragraph, *, first_line: bool):
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(24 if first_line else 0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def set_cell_fill(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = BORDER, size: str = "4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:cantSplit")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def add_page_field(paragraph):
    paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, result, end):
        run = paragraph.add_run()
        run._r.append(item)
    paragraph.add_run(" 页")
    for run in paragraph.runs:
        format_run(run, size=9)


def add_numbering_definition(doc: Document, *, decimal: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if decimal else "bullet")
    level.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1．" if decimal else "•")
    level.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "480")
    ind.set(qn("w:hanging"), "240")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "480")
    spacing.set(qn("w:lineRule"), "exact")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:numPr"))
    if existing is not None:
        ppr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def add_inline(paragraph, text: str, *, bold=False):
    token_re = re.compile(r"(\*\*[^*]+\*\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            format_run(run, bold=bold)
        run = paragraph.add_run(match.group(0)[2:-2])
        format_run(run, bold=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        format_run(run, bold=bold)


def configure_document(doc: Document):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.4)
        section.gutter = Cm(0)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    normal = doc.styles["Normal"]
    format_style(normal, BODY_LATIN, BODY_CJK, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(24)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    list_style = doc.styles["List Paragraph"]
    format_style(list_style, BODY_LATIN, BODY_CJK, 12)
    list_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    list_style.paragraph_format.line_spacing = Pt(24)
    list_style.paragraph_format.first_line_indent = Pt(0)
    list_style.paragraph_format.space_before = Pt(0)
    list_style.paragraph_format.space_after = Pt(0)

    heading_styles = {
        "Heading 1": (H1_LATIN, H1_CJK, True),
        "Heading 2": (H2_LATIN, H2_CJK, False),
        "Heading 3": (BODY_LATIN, BODY_CJK, False),
        "Heading 4": (BODY_LATIN, BODY_CJK, False),
    }
    for name, (latin, cjk, bold) in heading_styles.items():
        style = doc.styles[name]
        format_style(style, latin, cjk, 14, bold=bold)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(24)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    format_style(caption, BODY_LATIN, BODY_CJK, 12, bold=False)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.line_spacing = Pt(24)
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)
    caption.paragraph_format.keep_with_next = True

    settings = doc.settings._element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def configure_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for paragraph in section.header.paragraphs:
        paragraph.clear()
    footer_p = section.footer.paragraphs[0]
    footer_p.clear()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_fixed_24(footer_p, first_line=False)
    add_page_field(footer_p)


def add_cover(doc: Document):
    spacer = doc.add_paragraph()
    spacer.add_run("\n\n\n")
    spacer.paragraph_format.space_after = Pt(0)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("合同研究报告 · 建设指导评审稿")
    format_run(run, H1_LATIN, H1_CJK, 12, bold=True)
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    title.paragraph_format.line_spacing = 1.08
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("基于轨道交通客运装备的\n内装模块分区识别与快速设计\n系统构建研究报告")
    format_run(run, H1_LATIN, H1_CJK, 27, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    run = subtitle.add_run(SYSTEM_NAME)
    format_run(run, H2_LATIN, H2_CJK, 14)

    meta = [
        f"文档编号：{DOC_ID}    文档版本：{DOC_VERSION}",
        "文档性质：平台构建前技术指导文件",
        "合同依据：技术开发合同（26版）及附件技术规格书",
        "委托单位：中车工业研究院有限公司",
        "承担单位：北京科技大学",
        COMPILED_DATE,
    ]
    for idx, line in enumerate(meta):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6 if idx < len(meta) - 1 else 0)
        run = p.add_run(line)
        format_run(run, BODY_LATIN, BODY_CJK, 11, bold=idx in {0, len(meta) - 1})

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_document(doc)
    configure_header_footer(body)


def add_table(doc: Document, rows: list[list[str]], weights: list[float] | None = None):
    cols = max(len(row) for row in rows)
    normalized = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    if weights is None:
        lengths = [max(len(row[col]) for row in normalized) for col in range(cols)]
        weights = [max(0.75, min(float(length), 5.5)) for length in lengths]
    widths = column_widths_from_weights(weights, TABLE_WIDTH_DXA)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        set_row_cant_split(row)
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                set_cell_fill(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_fixed_24(p, first_line=False)
            add_inline(p, normalized[row_idx][col_idx], bold=row_idx == 0)
    set_repeat_table_header(table.rows[0])
    return table


def add_callout(doc: Document, title: str, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.style = None
    apply_table_geometry(
        table,
        [TABLE_WIDTH_DXA],
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa={"top": 120, "bottom": 120, "start": 160, "end": 160},
    )
    set_table_borders(table, color="D7DCE2", size="5")
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_fill(cell, "F7F8FA")
    p = cell.paragraphs[0]
    set_fixed_24(p, first_line=False)
    run = p.add_run(title)
    format_run(run, H1_LATIN, H1_CJK, 12, bold=True)
    p2 = cell.add_paragraph()
    set_fixed_24(p2, first_line=False)
    add_inline(p2, text)


def extract_toc_entries(source: str) -> list[tuple[int, str, str]]:
    lines = source.splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1 :]
    entries = []
    for line in lines:
        match = re.match(r"^(#{1,3})\s+(.+)$", line.strip())
        if not match:
            continue
        level = len(match.group(1))
        text = match.group(2)
        entries.append((level, text, f"report_heading_{len(entries) + 1:03d}"))
    return entries


def add_internal_hyperlink(paragraph, text: str, anchor: str, *, bold=False):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = paragraph.add_run(text)
    format_run(run, size=12, bold=bold)
    run.font.color.rgb = BLACK
    run.font.underline = False
    paragraph._p.remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def add_heading_bookmark(paragraph, bookmark_name: str, bookmark_id: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(1, start)
    paragraph._p.append(end)


def add_clickable_toc(doc: Document, toc_entries: list[tuple[int, str, str]]):
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("目录")
    format_run(run, H1_LATIN, H1_CJK, 14, bold=True)

    for level, text, anchor in toc_entries:
        p = doc.add_paragraph()
        set_fixed_24(p, first_line=False)
        p.paragraph_format.left_indent = Pt(0 if level == 1 else 24 if level == 2 else 48)
        p.paragraph_format.right_indent = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.7),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )
        add_internal_hyperlink(p, text, anchor, bold=level == 1)
        tab = p.add_run("\t")
        format_run(tab)
        page = str(TOC_PAGE_MAP.get(text, "—"))
        add_internal_hyperlink(p, page, anchor, bold=level == 1)


def add_front_matter(doc: Document, toc_entries: list[tuple[int, str, str]]):
    h = doc.add_heading("文档控制", level=1)
    h.paragraph_format.space_before = Pt(0)
    add_table(
        doc,
        [
            ["控制项", "记录"],
            ["文档编号 / 版本", f"{DOC_ID} / {DOC_VERSION}"],
            ["文档状态", "建设指导评审稿"],
            ["编制日期", COMPILED_DATE],
            ["文档定位", "平台构建前的关键技术、总体架构、功能模块、技术路线和进度指导文件"],
            ["合同依据", "技术开发合同（26版）及附件技术规格书"],
            ["内容参考", "技术响应第四版及甲方关于报告定位、图文表达和建设指导性的补充意见"],
        ],
        weights=[1.45, 5.05],
    )
    h2 = doc.add_heading("修订记录", level=2)
    h2.paragraph_format.space_before = Pt(12)
    add_table(
        doc,
        [
            ["版本", "日期", "状态", "修订说明"],
            ["V1.1", "2026-09-03", "历史评审稿", "依据合同扩充研究内容、方法、指标和追踪矩阵"],
            ["V1.2", "2026-09-03", "历史格式稿", "按指定页面、字体、行距、缩进和标题编号统一排版"],
            [DOC_VERSION, "2026-09-03", "当前评审稿", "按甲方意见重构为平台建设指导文件，精简文字，增加统一编号图表和可点击目录"],
        ],
        weights=[0.7, 1.0, 1.2, 3.6],
    )
    add_callout(
        doc,
        "报告使用说明",
        "本报告用于统一平台建设目标、关键技术、总体架构、数据与功能边界、实施路线和验证方法。具体模型、参数、接口字段和部署资源在详细设计与现场确认阶段固化；涉及工程安全、尺寸、装配和法规的结论以专业设计、试验及审批结果为准。",
    )
    doc.add_page_break()
    add_clickable_toc(doc, toc_entries)
    doc.add_page_break()


def add_picture(doc: Document, image_path: Path, caption_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    # A drawing is taller than a text line. Applying the body style's exact
    # 24-point line box here clips the image in Word/LibreOffice.
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run()
    inline = run.add_picture(str(image_path), width=Inches(6.18))
    inline._inline.docPr.set("descr", caption_text)
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_with_next = False
    run = cap.add_run(caption_text)
    format_run(run)


def parse_markdown(doc: Document, source: str, toc_entries: list[tuple[int, str, str]]):
    lines = source.splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1 :]

    i = 0
    active_decimal = None
    heading_index = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for row_idx, table_line in enumerate(table_lines):
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if row_idx == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            active_decimal = None
            continue

        image_match = re.fullmatch(r"!\[(图\s*\d+\s+.+?)\]\((.+?)\)", stripped)
        if image_match:
            caption, rel_path = image_match.groups()
            add_picture(doc, (SOURCE_PATH.parent / rel_path).resolve(), caption)
            active_decimal = None
            i += 1
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            p = doc.add_heading(text, level=level)
            expected_level, expected_text, bookmark = toc_entries[heading_index]
            if expected_level != level or expected_text != text:
                raise ValueError(f"heading order mismatch: {text}")
            add_heading_bookmark(p, bookmark, heading_index + 1)
            heading_index += 1
            set_fixed_24(p, first_line=False)
            latin, cjk, bold = {
                1: (H1_LATIN, H1_CJK, True),
                2: (H2_LATIN, H2_CJK, False),
                3: (BODY_LATIN, BODY_CJK, False),
                4: (BODY_LATIN, BODY_CJK, False),
            }[level]
            for run in p.runs:
                format_run(run, latin, cjk, 14, bold=bold)
            active_decimal = None
            i += 1
            continue

        if re.fullmatch(r"表\s*\d+\s+.+", stripped):
            p = doc.add_paragraph(style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped)
            format_run(run)
            active_decimal = None
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            if active_decimal is None:
                active_decimal = add_numbering_definition(doc, decimal=True)
            p = doc.add_paragraph(style="List Paragraph")
            set_fixed_24(p, first_line=False)
            apply_num(p, active_decimal)
            add_inline(p, numbered.group(1))
            i += 1
            continue

        if re.match(r"^-\s+", stripped):
            active_decimal = None
            bullet = add_numbering_definition(doc, decimal=False)
            p = doc.add_paragraph(style="List Paragraph")
            set_fixed_24(p, first_line=False)
            apply_num(p, bullet)
            add_inline(p, re.sub(r"^-\s+", "", stripped))
            i += 1
            continue

        active_decimal = None
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if (
                not candidate
                or candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith("![")
                or re.fullmatch(r"表\s*\d+\s+.+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
                or candidate.startswith("- ")
            ):
                break
            paragraph_lines.append(candidate)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        set_fixed_24(p, first_line=True)
        add_inline(p, " ".join(paragraph_lines))


def fbox(draw, x, y, w, h, title, sub="", *, fill="F4F6F9", edge="2E74B5"):
    figure_base.draw_box(draw, (x, y), w, h, title, sub, fill=fill, edge=edge)


def farrow(draw, start, end, *, color="2E74B5", rad=0.0):
    figure_base.arrow(draw, start, end, color=color, rad=rad)


def save_figure(image, name: str):
    image.save(ASSETS_DIR / name, format="PNG", optimize=True)


def generate_figures():
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)

    image, draw = figure_base.setup_figure("从合同与技术标书到平台落地的指导链")
    items = [
        (0.03, "合同与甲方要求", "目标 · 范围\n成果 · 验收"),
        (0.22, "技术标书", "总体响应\n技术承诺"),
        (0.41, "建设指导报告", "关键技术 · 架构\n功能 · 路线 · 计划"),
        (0.62, "详细设计与实施", "数据 · 接口\n开发 · 集成"),
        (0.82, "测试与交付", "指标 · 证据\n评审 · 验收"),
    ]
    for x, title, sub in items:
        fbox(draw, x, 0.44, 0.15, 0.27, title, sub, fill="FFF0F2" if x in {0.03, 0.82} else "F4F6F9", edge=ACCENT if x in {0.03, 0.82} else "2E74B5")
    for x in (0.18, 0.37, 0.56, 0.77):
        farrow(draw, (x, 0.575), (x + 0.04, 0.575), color=ACCENT)
    draw.text(figure_base._point_pixels((0.5, 0.20)), "报告既承接投标承诺，也为需求、数据、设计、开发和验收提供统一依据", font=figure_base.figure_font(27), fill="#5E6975", anchor="mm")
    save_figure(image, "v2-figure-01-guidance-framework.png")

    image, draw = figure_base.setup_figure("项目上下文中的快速设计业务闭环")
    loop = [
        (0.03, "设计需求", "车型 · 区域\n目标 · 约束"),
        (0.18, "资料准备", "图像 · 图纸\n模型 · 文本"),
        (0.33, "分区识别", "候选 · 复核\n模块映射"),
        (0.48, "方案生成", "模块 · 参数\n形态 · CMF"),
        (0.63, "规则校核", "接口 · 边界\n协同 · 安全"),
        (0.78, "评审入库", "对比 · 评价\n版本 · 归档"),
    ]
    for x, title, sub in loop:
        fbox(draw, x, 0.48, 0.12, 0.25, title, sub, edge=ACCENT if x >= 0.63 else "2E74B5")
    for x in (0.15, 0.30, 0.45, 0.60, 0.75):
        farrow(draw, (x, 0.605), (x + 0.03, 0.605), color=ACCENT if x >= 0.60 else "2E74B5")
    fbox(draw, 0.34, 0.12, 0.28, 0.18, "成果检索与复用", "历史方案反向支持新项目，保留来源与版本关系", fill="FFF0F2", edge=ACCENT)
    farrow(draw, (0.84, 0.48), (0.58, 0.30), color=ACCENT)
    farrow(draw, (0.34, 0.21), (0.09, 0.48), color=ACCENT, rad=-0.16)
    save_figure(image, "v2-figure-02-business-loop.png")

    image, draw = figure_base.setup_figure("研究对象、关键技术、平台能力与验证证据四条主线")
    streams = [
        (0.05, "对象主线", "区域 · 四级模块\n属性 · 接口 · 版本"),
        (0.28, "技术主线", "多模态识别\n生成 · 规则 · 优化"),
        (0.51, "平台主线", "项目 · 资源 · 任务\n权限 · 成果 · 接口"),
        (0.74, "验证主线", "数据集 · 指标 · 场景\n测试 · 评审 · 交付"),
    ]
    for x, title, sub in streams:
        fbox(draw, x, 0.48, 0.18, 0.29, title, sub, fill="E8EEF5" if x in {0.05, 0.51} else "FFF0F2", edge="2E74B5" if x in {0.05, 0.51} else ACCENT)
        farrow(draw, (x + 0.09, 0.48), (x + 0.09, 0.31), color="1F4D78")
    fbox(draw, 0.27, 0.12, 0.46, 0.19, "统一建设蓝图", "需求可追踪 · 对象可表达 · 能力可替换 · 过程可复现 · 结果可验证", fill="F4F6F9", edge="1F4D78")
    save_figure(image, "v2-figure-03-research-streams.png")

    image, draw = figure_base.setup_figure("客运装备内装模块化分区快速设计平台四层架构")
    layers = [
        (0.66, "应用层", "设计工作台 · 分区识别 · 快速设计 · 整体布置 · CMF · 方案评审", "FFF0F2", ACCENT),
        (0.48, "核心能力层", "多模态理解 · 候选识别 · 参数生成 · 规则校核 · 检索评价 · 跨区域优化", "E8EEF5", "2E74B5"),
        (0.30, "平台支撑层", "身份权限 · 项目协作 · 资源版本 · 任务编排 · 通知审计 · 接口适配", "F4F6F9", "1F4D78"),
        (0.12, "数据基础层", "基础资料 · 模块库 · 材料库 · 规则参数库 · 成果库 · 运行与审计数据", "F4F6F9", "5E6975"),
    ]
    for y, title, sub, fill, edge in layers:
        fbox(draw, 0.08, y, 0.84, 0.13, title, sub, fill=fill, edge=edge)
    for y in (0.66, 0.48, 0.30):
        farrow(draw, (0.50, y), (0.50, y - 0.05), color="2E74B5")
    save_figure(image, "v2-figure-04-layered-architecture.png")

    image, draw = figure_base.setup_figure("平台服务边界与受控部署拓扑")
    fbox(draw, 0.03, 0.42, 0.14, 0.24, "用户访问区", "统一入口\n身份认证\n设计工作台", fill="FFF0F2", edge=ACCENT)
    fbox(draw, 0.22, 0.42, 0.17, 0.24, "平台服务区", "项目 · 资源\n任务 · 权限\n版本 · 审计", fill="E8EEF5")
    fbox(draw, 0.45, 0.42, 0.17, 0.24, "能力服务区", "识别 · 理解\n生成 · 检索\n优化 · 报告", fill="E8EEF5")
    fbox(draw, 0.68, 0.59, 0.14, 0.18, "数据基础设施", "业务数据\n文件与知识", fill="F4F6F9", edge="1F4D78")
    fbox(draw, 0.68, 0.31, 0.14, 0.18, "计算资源区", "通用计算\n专用计算", fill="F4F6F9", edge="1F4D78")
    fbox(draw, 0.86, 0.42, 0.11, 0.24, "既有产品平台", "主数据\n任务与成果\n受控交互", fill="FFF0F2", edge=ACCENT)
    for start, end in [((0.17, 0.54), (0.22, 0.54)), ((0.39, 0.54), (0.45, 0.54)), ((0.62, 0.58), (0.68, 0.68)), ((0.62, 0.48), (0.68, 0.40)), ((0.82, 0.54), (0.86, 0.54))]:
        farrow(draw, start, end, color=ACCENT if end[0] >= 0.86 else "2E74B5")
    draw.text(figure_base._point_pixels((0.50, 0.16)), "浏览器不直接访问模型、计算节点或私有存储；敏感配置只在服务端受控管理", font=figure_base.figure_font(25), fill="#5E6975", anchor="mm")
    save_figure(image, "v2-figure-05-deployment-boundary.png")

    image, draw = figure_base.setup_figure("从客室系统到零件的四级模块对象体系")
    levels = [
        (0.10, 0.67, 0.80, "系统级", "客室内装系统 · 车型 · 目标市场 · 总体风格"),
        (0.18, 0.50, 0.64, "子系统级", "侧墙 · 顶部 · 地板 · 端墙 · 座椅 · 车门周边"),
        (0.27, 0.33, 0.46, "组件级", "可配置与装配单元 · 参数 · 接口 · 适用区域"),
        (0.36, 0.16, 0.28, "零件级", "表面 · 材料 · 工艺 · 公差 · 维护信息"),
    ]
    for x, y, w, title, sub in levels:
        fbox(draw, x, y, w, 0.13, title, sub, fill="FFF0F2" if title in {"系统级", "零件级"} else "E8EEF5", edge=ACCENT if title in {"系统级", "零件级"} else "2E74B5")
        if y > 0.16:
            farrow(draw, (0.50, y), (0.50, y - 0.04), color="1F4D78")
    save_figure(image, "v2-figure-06-module-hierarchy.png")

    image, draw = figure_base.setup_figure("多模态资料到可用分区对象的识别流程")
    sources = [(0.03, "图像", "外观 · 边界"), (0.03, "图纸", "轮廓 · 尺寸"), (0.03, "三维", "几何 · 装配"), (0.03, "文本属性", "名称 · 规则")]
    ys = [0.70, 0.52, 0.34, 0.16]
    for (_, title, sub), y in zip(sources, ys):
        fbox(draw, 0.03, y, 0.14, 0.13, title, sub, fill="F4F6F9", edge="2E74B5")
        farrow(draw, (0.17, y + 0.065), (0.25, 0.53), color="2E74B5")
    fbox(draw, 0.25, 0.40, 0.16, 0.26, "预处理与特征提取", "质量检查 · 对齐\n视觉 · 几何\n语义 · 关系", fill="E8EEF5")
    fbox(draw, 0.47, 0.40, 0.14, 0.26, "多模态融合", "区域先验\n对象词典\n空间与规则", fill="E8EEF5")
    fbox(draw, 0.67, 0.56, 0.14, 0.18, "候选分区", "类别 · 边界\n置信度 · 证据", fill="FFF0F2", edge=ACCENT)
    fbox(draw, 0.67, 0.28, 0.14, 0.18, "规则与人工复核", "冲突检查\n确认与修订", fill="FFF0F2", edge=ACCENT)
    fbox(draw, 0.86, 0.40, 0.11, 0.26, "正式识别结果", "区域对象\n模块映射\n属性与版本", fill="F4F6F9", edge="1F4D78")
    for start, end in [((0.41, 0.53), (0.47, 0.53)), ((0.61, 0.56), (0.67, 0.65)), ((0.74, 0.56), (0.74, 0.46)), ((0.81, 0.37), (0.86, 0.53))]:
        farrow(draw, start, end, color=ACCENT if start[0] >= 0.61 else "2E74B5")
    save_figure(image, "v2-figure-07-multimodal-recognition.png")

    image, draw = figure_base.setup_figure("候选识别结果的人机协同分流")
    fbox(draw, 0.04, 0.43, 0.16, 0.26, "候选结果", "类别 · 边界\n置信度 · 规则状态", fill="E8EEF5")
    branches = [
        (0.30, 0.66, "高置信且无冲突", "快速确认", "EAF7EF", "2D7D46"),
        (0.30, 0.42, "中等置信或边界不稳", "重点复核与局部修订", "FFF7E6", "B7791F"),
        (0.30, 0.18, "低置信或违反硬约束", "重新标注或退回处理", "FFF0F2", ACCENT),
    ]
    for x, y, title, sub, fill, edge in branches:
        fbox(draw, x, y, 0.28, 0.16, title, sub, fill=fill, edge=edge)
        farrow(draw, (0.20, 0.56), (x, y + 0.08), color=edge)
    fbox(draw, 0.68, 0.43, 0.16, 0.26, "人工确认", "类别 · 边界\n实例 · 适用范围", fill="E8EEF5", edge="1F4D78")
    for y, _, _, _, edge in [(0.66, "", "", "", "2D7D46"), (0.42, "", "", "", "B7791F"), (0.18, "", "", "", ACCENT)]:
        farrow(draw, (0.58, y + 0.08), (0.68, 0.56), color=edge)
    fbox(draw, 0.88, 0.43, 0.09, 0.26, "版本化结果", "原始候选\n修订记录\n最终结果", fill="F4F6F9", edge=ACCENT)
    farrow(draw, (0.84, 0.56), (0.88, 0.56), color=ACCENT)
    save_figure(image, "v2-figure-08-human-review.png")

    image, draw = figure_base.setup_figure("多模态能力编排与工程约束")
    steps = [
        (0.03, "设计意图", "区域 · 目标\n参考 · 约束"),
        (0.19, "对象识别", "区域与模块\n候选定位"),
        (0.35, "知识检索", "术语 · 规则\n相似方案"),
        (0.51, "方案生成", "形态 · 材质\n色彩 · 纹理"),
        (0.67, "结构与规则检查", "边界 · 接口\n一致性"),
        (0.83, "人工评审", "比较 · 调整\n确认 · 入库"),
    ]
    for x, title, sub in steps:
        fbox(draw, x, 0.48, 0.12, 0.25, title, sub, fill="F4F6F9", edge=ACCENT if x >= 0.67 else "2E74B5")
    for x in (0.15, 0.31, 0.47, 0.63, 0.79):
        farrow(draw, (x, 0.605), (x + 0.04, 0.605), color=ACCENT if x >= 0.63 else "2E74B5")
    fbox(draw, 0.23, 0.14, 0.54, 0.18, "统一任务与版本记录", "输入资料 · 业务参数 · 规则版本 · 能力版本 · 过程状态 · 输出成果", fill="E8EEF5", edge="1F4D78")
    save_figure(image, "v2-figure-09-capability-orchestration.png")

    image, draw = figure_base.setup_figure("模块化、参数化、规则化的快速设计引擎")
    inputs = [(0.03, "模块库", "可选对象"), (0.03, "参数模板", "可调范围"), (0.03, "规则库", "允许条件"), (0.03, "参考资料", "风格与目标")]
    for (_, title, sub), y in zip(inputs, [0.70, 0.52, 0.34, 0.16]):
        fbox(draw, 0.03, y, 0.14, 0.13, title, sub, fill="F4F6F9", edge="2E74B5")
        farrow(draw, (0.17, y + 0.065), (0.28, 0.53), color="2E74B5")
    fbox(draw, 0.28, 0.38, 0.18, 0.30, "快速设计引擎", "模块组合\n参数驱动\n生成式探索\n方案版本", fill="E8EEF5", edge="1F4D78")
    fbox(draw, 0.54, 0.58, 0.16, 0.18, "硬约束过滤", "边界 · 接口\n干涉 · 安全", fill="FFF0F2", edge=ACCENT)
    fbox(draw, 0.54, 0.32, 0.16, 0.18, "软目标评价", "美观 · 协调\n成本 · 复用", fill="FFF0F2", edge=ACCENT)
    fbox(draw, 0.78, 0.45, 0.18, 0.24, "候选方案集", "对比 · 排序 · 调整\n评审 · 入库 · 复用", fill="F4F6F9", edge="1F4D78")
    farrow(draw, (0.46, 0.58), (0.54, 0.67), color=ACCENT)
    farrow(draw, (0.46, 0.47), (0.54, 0.41), color=ACCENT)
    farrow(draw, (0.70, 0.67), (0.78, 0.60), color=ACCENT)
    farrow(draw, (0.70, 0.41), (0.78, 0.53), color=ACCENT)
    farrow(draw, (0.84, 0.45), (0.39, 0.38), color=ACCENT, rad=0.23)
    save_figure(image, "v2-figure-10-rapid-design-engine.png")

    image, draw = figure_base.setup_figure("客室区域依赖图与协同优化")
    nodes = {
        "顶部": (0.43, 0.70), "侧墙": (0.20, 0.48), "端墙": (0.43, 0.48),
        "门区": (0.66, 0.48), "地板": (0.43, 0.24), "座椅": (0.20, 0.24),
    }
    for title, (x, y) in nodes.items():
        fbox(draw, x, y, 0.14, 0.13, title, "接口 · 风格 · 装配", fill="E8EEF5", edge="2E74B5")
    links = [("顶部", "侧墙"), ("顶部", "端墙"), ("顶部", "门区"), ("侧墙", "端墙"), ("侧墙", "地板"), ("侧墙", "座椅"), ("端墙", "地板"), ("门区", "地板"), ("地板", "座椅")]
    for a, b in links:
        ax, ay = nodes[a]; bx, by = nodes[b]
        farrow(draw, (ax + 0.07, ay + 0.065), (bx + 0.07, by + 0.065), color="9AA7B5")
    fbox(draw, 0.79, 0.35, 0.18, 0.32, "协同优化顺序", "1 硬约束过滤\n2 条件规则限制\n3 软目标排序\n4 人工评审确认", fill="FFF0F2", edge=ACCENT)
    farrow(draw, (0.57, 0.54), (0.79, 0.54), color=ACCENT)
    save_figure(image, "v2-figure-11-cross-region-optimization.png")

    image, draw = figure_base.setup_figure("平台功能模块地图")
    fbox(draw, 0.04, 0.58, 0.18, 0.20, "设计工作台", "项目上下文\n资料 · 任务 · 结果", fill="FFF0F2", edge=ACCENT)
    modules = [
        (0.28, 0.67, "分区识别", "候选 · 复核"), (0.47, 0.67, "快速设计", "模块 · 参数"),
        (0.66, 0.67, "整体布置", "组合 · 协同"), (0.28, 0.42, "形态与CMF", "材质 · 色彩"),
        (0.47, 0.42, "方案成果库", "评价 · 复用"), (0.66, 0.42, "任务与通知", "状态 · 异常"),
    ]
    for x, y, title, sub in modules:
        fbox(draw, x, y, 0.15, 0.16, title, sub, fill="E8EEF5", edge="2E74B5")
        farrow(draw, (0.22, 0.68), (x, y + 0.08), color="2E74B5")
    fbox(draw, 0.28, 0.15, 0.53, 0.15, "平台管理", "项目协作 · 资源版本 · 规则能力 · 用户权限 · 运行审计 · 接口管理", fill="F4F6F9", edge="1F4D78")
    fbox(draw, 0.85, 0.42, 0.12, 0.36, "角色", "设计人员\n项目负责人\n算法与数据\n系统管理", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-figure-12-function-map.png")

    image, draw = figure_base.setup_figure("平台数据资源与知识底座")
    data = [
        (0.04, 0.60, "客室基础资料", "车型 · 图纸\n模型 · 图像"),
        (0.22, 0.60, "内装模块库", "四级对象\n属性 · 接口"),
        (0.40, 0.60, "材料CMF库", "材料 · 色彩\n纹理 · 工艺"),
        (0.58, 0.60, "规则参数库", "约束 · 条件\n模板 · 版本"),
        (0.76, 0.60, "设计成果库", "方案 · 评价\n评审 · 复用"),
    ]
    for x, y, title, sub in data:
        fbox(draw, x, y, 0.15, 0.19, title, sub, fill="E8EEF5", edge="2E74B5")
    fbox(draw, 0.12, 0.32, 0.76, 0.16, "统一数据治理", "分类编码 · 项目归属 · 标签检索 · 版本关系 · 来源授权 · 质量检查 · 生命周期", fill="F4F6F9", edge="1F4D78")
    fbox(draw, 0.24, 0.12, 0.52, 0.14, "结构化对象与文件资料通过资源标识关联", "设计输入、任务过程、输出成果和评审结论形成完整链路", fill="FFF0F2", edge=ACCENT)
    for x in (0.115, 0.295, 0.475, 0.655, 0.835):
        farrow(draw, (x, 0.60), (x, 0.48), color="2E74B5")
    save_figure(image, "v2-figure-13-data-foundation.png")

    image, draw = figure_base.setup_figure("平台建设技术路线与三阶段成果")
    steps = [
        (0.02, "需求与数据", "范围 · 对象\n样本 · 接口"), (0.155, "关键技术", "识别 · 多模态\n设计 · 优化"),
        (0.29, "标准与规则", "模块 · 参数\n约束 · 评价"), (0.425, "平台构建", "功能 · 数据\n任务 · 权限"),
        (0.56, "能力与接口", "模型能力\n产品平台"), (0.695, "验证迭代", "技术 · 功能\n场景 · 安全"),
        (0.83, "成果固化", "系统 · 数据\n文档 · 证据"),
    ]
    for x, title, sub in steps:
        fbox(draw, x, 0.58, 0.12, 0.21, title, sub, fill="F4F6F9", edge="2E74B5")
    for x in (0.14, 0.275, 0.41, 0.545, 0.68, 0.815):
        farrow(draw, (x, 0.685), (x + 0.015, 0.685), color=ACCENT)
    milestones = [
        (0.06, "第一阶段", "2026-09-30前\n研究报告与总体设计"),
        (0.37, "第二阶段", "2026-12-31前\n平台构建与检测"),
        (0.68, "第三阶段", "2027-03-31前\n集成完善与交付"),
    ]
    for x, title, sub in milestones:
        fbox(draw, x, 0.18, 0.25, 0.22, title, sub, fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-figure-14-technical-roadmap.png")

    image, draw = figure_base.setup_figure("从关键技术到合同验收的分层验证体系")
    levels = [
        (0.08, 0.69, 0.84, "验收资料核查", "成果物、版本、测试证据、评审与整改闭环"),
        (0.15, 0.53, 0.70, "业务场景验证", "典型项目、完整设计流程、角色协作与异常恢复"),
        (0.22, 0.37, 0.56, "系统集成验证", "接口、状态一致、权限隔离、性能、安全与稳定运行"),
        (0.29, 0.21, 0.42, "模块功能验证", "输入输出、规则、任务、版本、检索和结果管理"),
        (0.36, 0.05, 0.28, "算法与规则验证", "精度、边界、修订成本、规则正确性和泛化能力"),
    ]
    for x, y, w, title, sub in levels:
        fbox(draw, x, y, w, 0.12, title, sub, fill="FFF0F2" if y in {0.69, 0.05} else "E8EEF5", edge=ACCENT if y in {0.69, 0.05} else "2E74B5")
    save_figure(image, "v2-figure-15-verification-system.png")


def validate_source(source: str):
    figures = [int(x) for x in re.findall(r"!\[图\s*(\d+)\s+", source)]
    tables = [int(x) for x in re.findall(r"^表\s*(\d+)\s+", source, flags=re.MULTILINE)]
    if figures != list(range(1, len(figures) + 1)):
        raise ValueError(f"figure numbering is not sequential: {figures}")
    if tables != list(range(1, len(tables) + 1)):
        raise ValueError(f"table numbering is not sequential: {tables}")


def build():
    source = SOURCE_PATH.read_text(encoding="utf-8")
    validate_source(source)
    toc_entries = extract_toc_entries(source)
    generate_figures()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_document(doc)
    configure_header_footer(doc.sections[0])
    props = doc.core_properties
    props.title = DOC_TITLE
    props.subject = "轨道交通客运装备内装模块分区识别与快速设计平台建设指导"
    props.author = "轨道客室智能设计平台项目组"
    props.keywords = "轨道交通, 客室内装, 模块分区, 多模态识别, 快速设计, 平台架构"
    props.comments = f"文档编号 {DOC_ID}；版本 {DOC_VERSION}；建设指导评审稿"

    add_cover(doc)
    add_front_matter(doc, toc_entries)
    parse_markdown(doc, source, toc_entries)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
