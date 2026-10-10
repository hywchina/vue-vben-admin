"""Check V2.5 targeted scope, content boundaries and rendered pagination.

Usage: python tools/audit_report_v2_5.py --pdf /path/to/rendered.pdf
This complements, but does not replace, visual inspection of every page.
"""
import argparse
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "deliverables/基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告_V2.5.docx"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pdf", type=Path, required=True)
    args = parser.parse_args()
    source = (ROOT / "source/REPORT_V2.5.md").read_text()
    source = source.split("---", 2)[2]
    plain_source = re.sub(r"!\[([^\]]+)\]\([^\)]+\)", r"\1", source)
    with ZipFile(DOCX) as package:
        assert package.testzip() is None
        xml = etree.fromstring(package.read("word/document.xml"))
        styles = etree.fromstring(package.read("word/styles.xml"))
        comments = etree.fromstring(package.read("word/comments.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
    doc = Document(DOCX)
    text = "\n".join(xml.xpath("//w:t/text()", namespaces=ns))
    svg_paths = [(ROOT / "source" / path).resolve().with_suffix(".svg") for path in re.findall(r"!\[[^\]]*\]\(([^)]+)\)", source)]
    svg_text = "\n".join("".join(etree.parse(str(p)).getroot().itertext()) for p in svg_paths if p.exists())
    banned = r"报价|价格|金额|预算|费用|付款|商务|报价单|投标|Git|pnpm|/api/v1|apps/|51c4095e|6b047f2e|5463e2ae"
    assert not re.search(banned, text + svg_text), "Unexpected business or code-detail content"
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    figs = [int(m.group(1)) for t in captions if (m := re.match(r"^图\s*(\d+)\s", t))]
    tables = [int(m.group(1)) for t in captions if (m := re.match(r"^表\s*(\d+)\s", t))]
    assert figs == list(range(1, 24)), figs
    assert tables == list(range(1, 14)), tables
    assert len(doc.tables) == 13 and len(doc.inline_shapes) == 23
    assert all(t.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None for t in doc.tables)
    assert all(n.get("descr") for n in xml.xpath("//wp:docPr", namespaces=ns))
    for section in doc.sections:
        expected = {"page_width": 21, "page_height": 29.7, "top_margin": 2.5, "bottom_margin": 2.5, "left_margin": 2.6, "right_margin": 2.4, "header_distance": 1.5, "footer_distance": 1.75, "gutter": 0}
        for attr, value in expected.items():
            assert abs(getattr(section, attr).cm - value) < 0.003, (attr, getattr(section, attr).cm)
    for name, size, cjk in [("Normal", 12, "宋体"), ("Heading 1", 14, "黑体"), ("Heading 2", 14, "楷体"), ("Heading 3", 14, "宋体"), ("Heading 4", 14, "宋体")]:
        style = doc.styles[name]
        assert style.font.size.pt == size, name
        assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == cjk, name
        assert style.paragraph_format.line_spacing.pt == 24, name
    assert not styles.xpath('//w:style[@w:styleId="Title"]//w:pBdr', namespaces=ns)
    bookmarks = set(xml.xpath("//w:bookmarkStart/@w:name", namespaces=ns))
    anchors = xml.xpath("//w:hyperlink/@w:anchor", namespaces=ns)
    expected_headings = len(re.findall(r"(?m)^#{1,4} ", source))
    assert len(bookmarks) == expected_headings and len(anchors) == 28
    assert set(anchors) <= bookmarks
    page_map = json.loads((ROOT / "source/TOC_V2.5.json").read_text())
    reader = PdfReader(args.pdf)
    def flat(items):
        for item in items:
            if isinstance(item, list):
                yield from flat(item)
            else:
                yield item.title, reader.get_destination_page_number(item) + 1
    headings = list(flat(reader.outline))
    assert len(headings) == expected_headings
    actual_map = {f"report_heading_{i:03d}": page for i, (_, page) in enumerate(headings, 1)}
    assert actual_map == page_map, {k: (page_map[k], v) for k, v in actual_map.items() if page_map[k] != v}
    assert len(reader.pages[1].get("/Annots", [])) == 28
    for paragraph in xml.xpath("//w:p[w:hyperlink]", namespaces=ns):
        links = paragraph.xpath("w:hyperlink", namespaces=ns)
        assert len(links) == 2
        assert int("".join(links[1].xpath(".//w:t/text()", namespaces=ns))) == page_map[links[1].get(qn("w:anchor"))]
    comment_ids = comments.xpath("//w:comment/@w:id", namespaces=ns)
    assert len(comment_ids) == 9 and len(set(comment_ids)) == 9
    assert sorted(comment_ids) == sorted(xml.xpath("//w:commentRangeStart/@w:id", namespaces=ns))
    assert sorted(comment_ids) == sorted(xml.xpath("//w:commentRangeEnd/@w:id", namespaces=ns))
    assert sorted(comment_ids) == sorted(xml.xpath("//w:commentReference/@w:id", namespaces=ns))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1[0] == "概述" and h1[-1] == "参考文献"
    assert not any(t.startswith(("摘要", "附录", "一、研究依据")) for t in h1)
    assert not any(p.text.startswith(("关键词：", "关键词:")) for p in doc.paragraphs)
    assert len(re.findall(r"(?m)^\[\d+\]", source)) == 8
    assert all(p.paragraph_format.page_break_before for p in doc.paragraphs if p.style.name == 'Heading 1')
    markers = re.findall(r'【待补图例 (M\d+)', source)
    assert markers == ['M01', 'M02', 'M03', 'M04', 'M05'], markers
    # Chapters 5–12 stay byte-identical after removing only figure-number changes.
    prior = (ROOT / 'source/REPORT_V2.4.md').read_text()
    def retained_body(value):
        body = value.split('# 五、平台功能模块与业务衔接', 1)[1].split('# 参考文献', 1)[0]
        return re.sub(r'图\s*\d+', '图 [N]', body)
    assert retained_body(source) == retained_body(prior), 'Out-of-scope chapter changes'
    ai_section = source.split("## （三）AI 聊天助手的图文辅助", 1)[1].split("## （四）", 1)[0]
    ai_section = re.sub(r"!\[.*?\]\(.*?\)", "", ai_section)
    stats = {"pages": len(reader.pages), "source_chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", plain_source)), "body_and_table_chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", text)), "assistant_section_chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", ai_section)), "figures": 23, "tables": 13, "heading_bookmarks": len(bookmarks), "toc_entries": 14, "toc_links": len(anchors), "toc_page_map_matches_render": True, "pricing_and_code_detail_matches": 0, "page_geometry_and_fonts": "pass", "zip_integrity": "pass", "native_revision_comments": 9, "technical_references": 8, "removed_abstract_keywords_and_appendices": True, "all_major_chapters_start_new_page": True, "pending_figure_markers": markers, "chapters_5_to_12_unchanged_except_figure_numbers": True}
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
