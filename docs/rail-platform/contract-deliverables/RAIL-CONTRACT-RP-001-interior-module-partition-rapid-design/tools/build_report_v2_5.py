#!/usr/bin/env python3
"""Build the targeted September meeting revision, preserving earlier artifacts."""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

import build_report_v2_3 as layout

BASE = Path(__file__).resolve().parents[1]
SOURCE = BASE / 'source' / 'REPORT_V2.5.md'
OUTPUT = BASE / 'deliverables' / '基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告_V2.5.docx'


def cover(doc):
    for _ in range(4):
        p = doc.add_paragraph(); layout.set_fixed_24(p, first_line=False)
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(36)
    p.paragraph_format.space_after = Pt(24)
    layout.format_style(doc.styles['Title'], layout.H1_LATIN, layout.H1_CJK, 24, bold=True)
    for element in (doc.styles['Title'].element, p._p):
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    layout.format_run(p.add_run('基于轨道交通客运装备的\n内装模块分区识别与快速设计\n系统构建研究报告'), layout.H1_LATIN, layout.H1_CJK, 24, bold=True)
    for value in ['平台构建技术指导', '', '文档版本  V2.5', '委托单位  中车工业研究院有限公司', '承担单位  北京科技大学', '', '2026年10月9日']:
        p = doc.add_paragraph(); layout.set_fixed_24(p, first_line=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        layout.format_run(p.add_run(value), size=12)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    layout.configure_document(doc); layout.configure_header_footer(section)
    for p in doc.sections[0].footer.paragraphs: p.clear()


def build():
    parser = argparse.ArgumentParser()
    parser.add_argument('--page-map', type=Path, default=SOURCE.with_name('TOC_V2.5.json'))
    parser.add_argument('--out', type=Path, default=OUTPUT)
    args = parser.parse_args()
    layout.SOURCE_PATH = SOURCE
    layout.TOC_PAGE_MAP = json.loads(args.page_map.read_text()) if args.page_map.exists() else {}
    source = SOURCE.read_text(encoding='utf-8')
    entries = layout.extract_toc_entries(source)
    doc = Document(); layout.configure_document(doc)
    doc.core_properties.title = layout.DOC_TITLE
    doc.core_properties.subject = '内装模块划分、识别方法、快速设计流程及平台构建方案'
    doc.core_properties.author = '轨道客室智能设计平台项目组'
    doc.core_properties.comments = 'RAIL-CONTRACT-RP-001 V2.5'
    cover(doc); layout.add_front_matter(doc, entries)
    layout.parse_markdown(doc, source, entries)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == 'Heading 1':
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.space_before = Pt(0)
        if paragraph.text.startswith('【待补图例'):
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string('A82020')
    # Use the native Word comments API so local revision explanations stay at
    # their relevant passages; original reviewer comments are separately archived.
    notes = [
        ('一、出口客运装备内装设计现状与研究方法', 'S01：所有章级标题另起一页，沿用既定字体、行距与页边距。'),
        ('出口客运装备内装设计需要兼顾产品族共性', 'S02：概述收束为背景、研究目的及形成的平台建设依据。'),
        ('项目面向中车出口产品平台开展研究', 'S03：补充出口车辆和客室特点、实际设计流程的分析框架；企业现状图及流程图就地标记 M01、M02 待补。'),
        ('（四）人工智能相关技术原理与方法选择', 'S04：合并原第一章 AI 分析与第二章方法内容；补充 ComfyUI、ControlNet、FLUX.2 原理图和功能说明，区分候选技术与已用技术。'),
        ('传统模块化划分首先确定', 'S05：保留传统模块划分方法，增加客室组成概念图；车型级工程拆解图仍以 M03 标注待补。'),
        ('颜色分区设计从客室原图开始', 'S06：第三章按独立工作流实践叙述，补充同源输入和两次既有生成结果，明确未满足指令的问题；原生标记图 M05 与图文理解案例 M04 待补。'),
        ('平台架构承接前文', 'S07：保持现状需求—模块及 AI 方法—快速设计实践—平台构建的章序，不删除第一章。'),
        ('应用交互采用Vue 3', 'S08：五层总图及表 6 明确实际 Vue 3、Node.js/Nitro/H3、PostgreSQL、MinIO/S3 等技术，不将专家举例的 Java/MySQL 写成事实。'),
        ('五、平台功能模块与业务衔接', 'S09：第五至十二章只随新增图机械调整图号，原有正文和结构保留；未改写后半部为其他主题。'),
    ]
    for anchor, text in notes:
        matches = [p for p in doc.paragraphs if p.text.startswith(anchor)
                   and (not anchor.startswith(('一、', '五、')) or p.style.name == 'Heading 1')]
        if len(matches) != 1: raise ValueError(f'Comment anchor not unique: {anchor}')
        doc.add_comment(matches[0].runs, text=text, author='修订说明', initials='RP')
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(args.out)


if __name__ == '__main__':
    build()
