#!/usr/bin/env python3
"""Build the formal V2.1 platform-construction guidance report."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

import build_report_docx as figure_base


SKILL_SCRIPTS = Path(
    "/home/huyanwei/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.826.12353/skills/documents/scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


BASE_DIR = Path(__file__).resolve().parents[1]
SOURCE_PATH = BASE_DIR / "source" / "REPORT_V2.1.md"
ASSETS_DIR = BASE_DIR / "assets"
OUTPUT_PATH = (
    BASE_DIR
    / "deliverables"
    / "基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告_V2.1.docx"
)

DOC_ID = "RAIL-CONTRACT-RP-001"
DOC_VERSION = "V2.1"
DOC_TITLE = "基于轨道交通客运装备的内装模块分区识别与快速设计系统构建研究报告"
SYSTEM_NAME = "客运装备内装模块化分区快速设计平台"
COMPILED_DATE = "2026年9月3日"

BODY_LATIN = "SimSun"
BODY_CJK = "宋体"
H1_LATIN = "SimHei"
H1_CJK = "黑体"
H2_LATIN = "KaiTi"
H2_CJK = "楷体"

BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 98, 108)
HEADER_FILL = "EEF1F5"
BORDER = "C9D1DA"
ACCENT = "B91C32"

CONTENT_WIDTH_DXA = 9072
TABLE_INDENT_DXA = 100
TABLE_WIDTH_DXA = CONTENT_WIDTH_DXA - TABLE_INDENT_DXA
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 100, "end": 100}

# Filled after the first deterministic render. The visible values make the
# table of contents useful even before Word refreshes fields; every entry also
# carries an internal hyperlink to its heading bookmark.
TOC_PAGE_MAP: dict[str, int] = {
    "摘要": 6,
    "一、报告定位与研究依据": 6,
    "（一）报告定位": 6,
    "（二）建设依据": 8,
    "（三）研究范围": 8,
    "（四）研究方法": 8,
    "二、建设背景与总体目标": 9,
    "（一）业务背景": 9,
    "（二）总体目标": 9,
    "（三）建设原则": 10,
    "三、总体研究框架与关键问题": 11,
    "（一）研究内容体系": 11,
    "（二）需要解决的关键问题": 12,
    "（三）研究成果形态": 12,
    "四、平台总体架构设计": 13,
    "（一）分层架构": 13,
    "（二）服务边界": 14,
    "（三）统一对象和业务协议": 15,
    "（四）部署与扩展架构": 15,
    "五、内装模块分区识别关键技术": 16,
    "（一）区域与模块对象模型": 16,
    "（二）多模态输入与预处理": 18,
    "（三）候选识别与多模态融合": 19,
    "（四）人机协同复核": 20,
    "（五）识别数据集与评价指标": 21,
    "六、多模态应用关键技术": 22,
    "（一）多模态能力定位": 22,
    "（二）语义对齐与知识增强": 22,
    "（三）能力编排": 23,
    "（四）模型选型与评测": 24,
    "七、快速设计系统关键技术": 24,
    "（一）模块化、参数化与规则化": 24,
    "（二）快速设计流程": 25,
    "（三）规则与约束体系": 26,
    "（四）跨区域协同优化": 27,
    "（五）方案检索、评价与复用": 27,
    "八、平台功能模块设计": 28,
    "（一）功能总体布局": 28,
    "（二）典型业务场景": 29,
    "（三）角色与权限": 30,
    "九、数据资源与知识底座": 30,
    "（一）数据资源体系": 30,
    "（二）数据治理流程": 31,
    "（三）数据质量控制": 32,
    "十、系统集成与接口设计": 32,
    "（一）与既有产品平台的交互": 32,
    "（二）外部智能能力接入": 33,
    "（三）接口验证": 33,
    "十一、安全、质量与运行保障": 33,
    "（一）安全与保密": 33,
    "（二）可追溯设计": 33,
    "（三）运行保障": 34,
    "十二、实施技术路线与进度计划": 34,
    "（一）总体实施路线": 34,
    "（二）阶段进度计划": 35,
    "（三）近期工作分解": 35,
    "（四）项目组织与沟通": 36,
    "十三、验证方法与验收建议": 36,
    "（一）分层验证方法": 36,
    "（二）典型验收场景": 37,
    "（三）质量指标处理原则": 37,
    "十四、主要风险与控制措施": 38,
    "（一）风险管理机制": 38,
    "十五、结论与建设建议": 39,
    "附录 A 甲方关注事项与报告章节对应表": 39,
    "附录 B 后续详细设计输入清单": 40,
}


def set_rfonts(target, latin: str, east_asia: str):
    rpr = target._element.get_or_add_rPr() if hasattr(target, "_element") else target
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def format_run(run, latin: str = BODY_LATIN, east_asia: str = BODY_CJK, size: float = 12, *, bold=None):
    run.font.name = latin
    set_rfonts(run, latin, east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def format_style(style, latin: str, east_asia: str, size: float, *, bold=None):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    if bold is not None:
        style.font.bold = bold
    set_rfonts(style, latin, east_asia)


def set_fixed_24(paragraph, *, first_line: bool):
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(24 if first_line else 0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def set_cell_fill(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = BORDER, size: str = "4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:cantSplit")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def add_page_field(paragraph):
    paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, result, end):
        run = paragraph.add_run()
        run._r.append(item)
    paragraph.add_run(" 页")
    for run in paragraph.runs:
        format_run(run, size=9)


def add_numbering_definition(doc: Document, *, decimal: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if decimal else "bullet")
    level.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1．" if decimal else "•")
    level.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "480")
    ind.set(qn("w:hanging"), "240")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "480")
    spacing.set(qn("w:lineRule"), "exact")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:numPr"))
    if existing is not None:
        ppr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def add_inline(paragraph, text: str, *, bold=False):
    token_re = re.compile(r"(\*\*[^*]+\*\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            format_run(run, bold=bold)
        run = paragraph.add_run(match.group(0)[2:-2])
        format_run(run, bold=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        format_run(run, bold=bold)


def configure_document(doc: Document):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.4)
        section.gutter = Cm(0)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    normal = doc.styles["Normal"]
    format_style(normal, BODY_LATIN, BODY_CJK, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(24)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    list_style = doc.styles["List Paragraph"]
    format_style(list_style, BODY_LATIN, BODY_CJK, 12)
    list_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    list_style.paragraph_format.line_spacing = Pt(24)
    list_style.paragraph_format.first_line_indent = Pt(0)
    list_style.paragraph_format.space_before = Pt(0)
    list_style.paragraph_format.space_after = Pt(0)

    heading_styles = {
        "Heading 1": (H1_LATIN, H1_CJK, True),
        "Heading 2": (H2_LATIN, H2_CJK, False),
        "Heading 3": (BODY_LATIN, BODY_CJK, False),
        "Heading 4": (BODY_LATIN, BODY_CJK, False),
    }
    for name, (latin, cjk, bold) in heading_styles.items():
        style = doc.styles[name]
        format_style(style, latin, cjk, 14, bold=bold)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(24)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    format_style(caption, BODY_LATIN, BODY_CJK, 12, bold=False)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.line_spacing = Pt(24)
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)
    caption.paragraph_format.keep_with_next = True

    settings = doc.settings._element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def configure_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for paragraph in section.header.paragraphs:
        paragraph.clear()
    footer_p = section.footer.paragraphs[0]
    footer_p.clear()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_fixed_24(footer_p, first_line=False)
    add_page_field(footer_p)


def add_cover(doc: Document):
    spacer = doc.add_paragraph()
    spacer.add_run("\n\n\n")
    spacer.paragraph_format.space_after = Pt(0)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("合同研究报告 · 建设指导评审稿")
    format_run(run, H1_LATIN, H1_CJK, 12, bold=True)
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    title.paragraph_format.line_spacing = 1.08
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("基于轨道交通客运装备的\n内装模块分区识别与快速设计\n系统构建研究报告")
    format_run(run, H1_LATIN, H1_CJK, 27, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    run = subtitle.add_run(SYSTEM_NAME)
    format_run(run, H2_LATIN, H2_CJK, 14)

    meta = [
        f"文档编号：{DOC_ID}    文档版本：{DOC_VERSION}",
        "文档性质：平台构建前技术指导文件",
        "合同依据：技术开发合同（26版）及附件技术规格书",
        "委托单位：中车工业研究院有限公司",
        "承担单位：北京科技大学",
        COMPILED_DATE,
    ]
    for idx, line in enumerate(meta):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6 if idx < len(meta) - 1 else 0)
        run = p.add_run(line)
        format_run(run, BODY_LATIN, BODY_CJK, 11, bold=idx in {0, len(meta) - 1})

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_document(doc)
    configure_header_footer(body)


def add_table(doc: Document, rows: list[list[str]], weights: list[float] | None = None):
    cols = max(len(row) for row in rows)
    normalized = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    if weights is None:
        lengths = [max(len(row[col]) for row in normalized) for col in range(cols)]
        weights = [max(0.75, min(float(length), 5.5)) for length in lengths]
    widths = column_widths_from_weights(weights, TABLE_WIDTH_DXA)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        set_row_cant_split(row)
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                set_cell_fill(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_fixed_24(p, first_line=False)
            add_inline(p, normalized[row_idx][col_idx], bold=row_idx == 0)
    set_repeat_table_header(table.rows[0])
    return table


def add_callout(doc: Document, title: str, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.style = None
    apply_table_geometry(
        table,
        [TABLE_WIDTH_DXA],
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa={"top": 120, "bottom": 120, "start": 160, "end": 160},
    )
    set_table_borders(table, color="D7DCE2", size="5")
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_fill(cell, "F7F8FA")
    p = cell.paragraphs[0]
    set_fixed_24(p, first_line=False)
    run = p.add_run(title)
    format_run(run, H1_LATIN, H1_CJK, 12, bold=True)
    p2 = cell.add_paragraph()
    set_fixed_24(p2, first_line=False)
    add_inline(p2, text)


def extract_toc_entries(source: str) -> list[tuple[int, str, str]]:
    lines = source.splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1 :]
    entries = []
    for line in lines:
        match = re.match(r"^(#{1,3})\s+(.+)$", line.strip())
        if not match:
            continue
        level = len(match.group(1))
        text = match.group(2)
        entries.append((level, text, f"report_heading_{len(entries) + 1:03d}"))
    return entries


def add_internal_hyperlink(paragraph, text: str, anchor: str, *, bold=False):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = paragraph.add_run(text)
    format_run(run, size=12, bold=bold)
    run.font.color.rgb = BLACK
    run.font.underline = False
    paragraph._p.remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def add_heading_bookmark(paragraph, bookmark_name: str, bookmark_id: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(1, start)
    paragraph._p.append(end)


def add_clickable_toc(doc: Document, toc_entries: list[tuple[int, str, str]]):
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("目录")
    format_run(run, H1_LATIN, H1_CJK, 14, bold=True)

    for level, text, anchor in toc_entries:
        p = doc.add_paragraph()
        set_fixed_24(p, first_line=False)
        p.paragraph_format.left_indent = Pt(0 if level == 1 else 24 if level == 2 else 48)
        p.paragraph_format.right_indent = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.7),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )
        add_internal_hyperlink(p, text, anchor, bold=level == 1)
        tab = p.add_run("\t")
        format_run(tab)
        page = str(TOC_PAGE_MAP.get(text, "—"))
        add_internal_hyperlink(p, page, anchor, bold=level == 1)


def add_front_matter(doc: Document, toc_entries: list[tuple[int, str, str]]):
    h = doc.add_heading("文档控制", level=1)
    h.paragraph_format.space_before = Pt(0)
    add_table(
        doc,
        [
            ["控制项", "记录"],
            ["文档编号 / 版本", f"{DOC_ID} / {DOC_VERSION}"],
            ["文档状态", "建设指导评审稿"],
            ["编制日期", COMPILED_DATE],
            ["文档定位", "平台构建前的关键技术、总体架构、功能模块、技术路线和进度指导文件"],
            ["合同依据", "技术开发合同（26版）及附件技术规格书"],
            ["内容参考", "技术响应第四版及甲方关于报告定位、图文表达和建设指导性的补充意见"],
        ],
        weights=[1.45, 5.05],
    )
    h2 = doc.add_heading("修订记录", level=2)
    h2.paragraph_format.space_before = Pt(12)
    add_table(
        doc,
        [
            ["版本", "日期", "状态", "修订说明"],
            ["V1.x", "2026-09-03", "历史评审稿", "合同研究内容扩充与指定格式统一"],
            ["V2.0", "2026-09-03", "历史指导稿", "重构为平台构建前建设指导文件"],
            [DOC_VERSION, "2026-09-03", "当前评审稿", "重绘技术图、清理无关内容并控制篇幅"],
        ],
        weights=[0.7, 1.0, 1.2, 3.6],
    )
    add_callout(
        doc,
        "报告使用说明",
        "本报告用于统一平台建设目标、关键技术、总体架构、功能边界、实施路线和验证方法；具体参数及工程结论在详细设计、试验和审批阶段固化。",
    )
    doc.add_page_break()
    add_clickable_toc(doc, toc_entries)
    doc.add_page_break()


def add_picture(doc: Document, image_path: Path, caption_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    # A drawing is taller than a text line. Applying the body style's exact
    # 24-point line box here clips the image in Word/LibreOffice.
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run()
    inline = run.add_picture(str(image_path), width=Inches(6.18))
    inline._inline.docPr.set("descr", caption_text)
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_with_next = False
    run = cap.add_run(caption_text)
    format_run(run)


def parse_markdown(doc: Document, source: str, toc_entries: list[tuple[int, str, str]]):
    lines = source.splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1 :]

    i = 0
    active_decimal = None
    heading_index = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for row_idx, table_line in enumerate(table_lines):
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if row_idx == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            active_decimal = None
            continue

        image_match = re.fullmatch(r"!\[(图\s*\d+\s+.+?)\]\((.+?)\)", stripped)
        if image_match:
            caption, rel_path = image_match.groups()
            add_picture(doc, (SOURCE_PATH.parent / rel_path).resolve(), caption)
            active_decimal = None
            i += 1
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            p = doc.add_heading(text, level=level)
            expected_level, expected_text, bookmark = toc_entries[heading_index]
            if expected_level != level or expected_text != text:
                raise ValueError(f"heading order mismatch: {text}")
            add_heading_bookmark(p, bookmark, heading_index + 1)
            heading_index += 1
            set_fixed_24(p, first_line=False)
            latin, cjk, bold = {
                1: (H1_LATIN, H1_CJK, True),
                2: (H2_LATIN, H2_CJK, False),
                3: (BODY_LATIN, BODY_CJK, False),
                4: (BODY_LATIN, BODY_CJK, False),
            }[level]
            for run in p.runs:
                format_run(run, latin, cjk, 14, bold=bold)
            active_decimal = None
            i += 1
            continue

        if re.fullmatch(r"表\s*\d+\s+.+", stripped):
            p = doc.add_paragraph(style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped)
            format_run(run)
            active_decimal = None
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            if active_decimal is None:
                active_decimal = add_numbering_definition(doc, decimal=True)
            p = doc.add_paragraph(style="List Paragraph")
            set_fixed_24(p, first_line=False)
            apply_num(p, active_decimal)
            add_inline(p, numbered.group(1))
            i += 1
            continue

        if re.match(r"^-\s+", stripped):
            active_decimal = None
            bullet = add_numbering_definition(doc, decimal=False)
            p = doc.add_paragraph(style="List Paragraph")
            set_fixed_24(p, first_line=False)
            apply_num(p, bullet)
            add_inline(p, re.sub(r"^-\s+", "", stripped))
            i += 1
            continue

        active_decimal = None
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if (
                not candidate
                or candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith("![")
                or re.fullmatch(r"表\s*\d+\s+.+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
                or candidate.startswith("- ")
            ):
                break
            paragraph_lines.append(candidate)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        set_fixed_24(p, first_line=True)
        add_inline(p, " ".join(paragraph_lines))


def fbox(draw, x, y, w, h, title, sub="", *, fill="F4F6F9", edge="2E74B5"):
    figure_base.draw_box(draw, (x, y), w, h, title, sub, fill=fill, edge=edge)


def farrow(draw, start, end, *, color="2E74B5", rad=0.0):
    figure_base.arrow(draw, start, end, color=color, rad=rad)


def fmini(draw, x, y, w, h, title, sub="", *, fill="F4F6F9", edge="2E74B5"):
    """Compact figure card for denser technical diagrams."""
    box = figure_base._box_pixels((x, y), w, h)
    draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{edge}", width=4)
    title_center = figure_base._point_pixels((x + w / 2, y + h * (0.64 if sub else 0.50)))
    draw.multiline_text(
        title_center,
        title,
        font=figure_base.figure_font(28),
        fill="#17375E",
        anchor="mm",
        align="center",
        spacing=4,
    )
    if sub:
        subtitle_center = figure_base._point_pixels((x + w / 2, y + h * 0.25))
        draw.multiline_text(
            subtitle_center,
            sub,
            font=figure_base.figure_font(20),
            fill="#5E6975",
            anchor="mm",
            align="center",
            spacing=4,
        )


def fortho(draw, points, *, color="2E74B5", arrow=True, width=5):
    """Draw a right-angle connector; callers route points to avoid crossings."""
    pixels = [figure_base._point_pixels(point) for point in points]
    draw.line(pixels, fill=f"#{color}", width=width, joint="curve")
    if arrow and len(points) >= 2:
        figure_base.arrow(draw, points[-2], points[-1], color=color)


def fsection_label(draw, x, y, text, *, color="17375E"):
    draw.text(
        figure_base._point_pixels((x, y)),
        text,
        font=figure_base.figure_font(22),
        fill=f"#{color}",
        anchor="lm",
    )


def save_figure(image, name: str):
    image.save(ASSETS_DIR / name, format="PNG", optimize=True)


def generate_figures():
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)

    image, draw = figure_base.setup_figure("合同要求向平台建设任务的分解与闭环")
    for x, title, sub in [
        (0.04, "合同与技术规格", "目标 · 范围 · 成果 · 验收"),
        (0.37, "技术响应文件", "总体方案 · 技术承诺 · 实施边界"),
        (0.70, "甲方补充意见", "报告定位 · 图文表达 · 落地指导"),
    ]:
        fmini(draw, x, 0.72, 0.26, 0.13, title, sub, fill="FFF0F2", edge=ACCENT)
        farrow(draw, (x + 0.13, 0.72), (x + 0.13, 0.64), color=ACCENT)
    fmini(draw, 0.04, 0.52, 0.92, 0.12, "建设边界与研究问题统一化", "对象标准、关键技术、业务流程、平台架构、数据条件、实施计划、验证证据", fill="E8EEF5", edge="1F4D78")
    columns = [
        (0.04, "对象与数据研究", "区域体系\n四级模块\n数据与知识"),
        (0.275, "关键技术研究", "分区识别\n多模态应用\n快速设计"),
        (0.51, "平台方案设计", "功能架构\n服务边界\n接口与安全"),
        (0.745, "实施验证设计", "阶段任务\n质量指标\n评审与交付"),
    ]
    for x, title, sub in columns:
        fmini(draw, x, 0.28, 0.205, 0.18, title, sub, fill="F4F6F9")
        farrow(draw, (x + 0.1025, 0.52), (x + 0.1025, 0.46))
        farrow(draw, (x + 0.1025, 0.28), (x + 0.1025, 0.20), color=ACCENT)
    outputs = ["标准与数据清单", "技术原型与规则", "总体设计与接口", "计划、指标与证据"]
    for (x, _, _), title in zip(columns, outputs):
        fmini(draw, x, 0.08, 0.205, 0.12, title, "可评审 · 可实施 · 可追踪", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-01-guidance-framework.png")

    image, draw = figure_base.setup_figure("项目上下文中的快速设计业务闭环与对象贯通")
    steps = [
        (0.025, "需求解析", "车型·市场"), (0.16, "资料治理", "图像·图纸"),
        (0.295, "分区识别", "候选·复核"), (0.43, "模块映射", "对象·参数"),
        (0.565, "方案生成", "形态·CMF"), (0.70, "协同校核", "接口·规则"),
        (0.835, "评审归档", "版本·复用"),
    ]
    for x, title, sub in steps:
        fmini(draw, x, 0.62, 0.115, 0.16, title, sub, fill="FFF0F2" if x in {0.025, 0.835} else "E8EEF5", edge=ACCENT if x in {0.025, 0.835} else "2E74B5")
    for x in (0.14, 0.275, 0.41, 0.545, 0.68, 0.815):
        farrow(draw, (x, 0.70), (x + 0.02, 0.70), color=ACCENT)
    fsection_label(draw, 0.03, 0.53, "贯通对象")
    objects = [
        (0.03, "项目上下文", "成员·范围·阶段"), (0.22, "区域与模块", "层级·属性·接口"),
        (0.41, "任务与过程", "输入·状态·证据"), (0.60, "方案与版本", "参数·规则·评价"),
        (0.79, "成果与知识", "血缘·检索·复用"),
    ]
    for x, title, sub in objects:
        fmini(draw, x, 0.31, 0.16, 0.14, title, sub, fill="F4F6F9", edge="1F4D78")
    for x in (0.11, 0.30, 0.49, 0.68, 0.87):
        farrow(draw, (x, 0.62), (x, 0.45), color="1F4D78")
    fmini(draw, 0.16, 0.09, 0.68, 0.12, "反馈与复用通道", "评审结论、修订记录和优选方案反向更新对象、规则、模板与知识", fill="FFF0F2", edge=ACCENT)
    fortho(draw, [(0.90, 0.31), (0.90, 0.23), (0.84, 0.23), (0.84, 0.15)], color=ACCENT)
    fortho(draw, [(0.16, 0.15), (0.08, 0.15), (0.08, 0.31)], color=ACCENT)
    save_figure(image, "v2-1-figure-02-business-loop.png")

    image, draw = figure_base.setup_figure("研究对象、研究方法与阶段成果矩阵")
    headers = [(0.22, "对象与标准"), (0.405, "关键技术"), (0.59, "平台与业务"), (0.775, "验证与实施")]
    for x, title in headers:
        fmini(draw, x, 0.75, 0.17, 0.10, title, fill="FFF0F2", edge=ACCENT)
    rows = [
        (0.58, "核心问题", ["区域和四级模块\n如何统一表达", "多模态识别与设计\n如何稳定协同", "能力如何进入真实\n项目业务闭环", "成果如何量化\n检查与追踪"]),
        (0.38, "研究方法", ["领域建模\n资料分析", "对照试验\n人机协同", "系统工程\n场景推演", "分层测试\n证据评审"]),
        (0.18, "阶段成果", ["编码、字典\n数据清单", "模型流程\n规则与指标", "架构、功能\n接口与原型", "计划、用例\n报告与清单"]),
    ]
    for y, label, cells in rows:
        fmini(draw, 0.04, y, 0.15, 0.15, label, "问题—方法—成果", fill="E8EEF5", edge="1F4D78")
        for (x, _), text_value in zip(headers, cells):
            fmini(draw, x, y, 0.17, 0.15, text_value, fill="F4F6F9", edge="9AA7B5")
    draw.text(figure_base._point_pixels((0.50, 0.08)), "四条研究主线共享同一需求基线、对象标识、版本关系和验证证据", font=figure_base.figure_font(24), fill="#5E6975", anchor="mm")
    save_figure(image, "v2-1-figure-03-research-streams.png")

    image, draw = figure_base.setup_figure("客运装备内装模块化分区快速设计平台总体技术架构")
    fmini(draw, 0.015, 0.13, 0.105, 0.69, "标准与治理", "模块标准\n数据规范\n接口规范\n质量基线\n变更控制\n知识产权", fill="FFF0F2", edge=ACCENT)
    fmini(draw, 0.88, 0.13, 0.105, 0.69, "安全与运行", "身份权限\n项目隔离\n配置保密\n审计追踪\n监控告警\n备份恢复", fill="FFF0F2", edge=ACCENT)
    layers = [
        (0.72, "业务访问层", "设计工作台｜管理门户｜评审验收｜产品平台协同｜移动/大屏展示", "FFF0F2", ACCENT),
        (0.595, "业务应用层", "项目协作｜分区识别｜模块化设计｜整体布置｜形态与CMF｜方案评审", "E8EEF5", "2E74B5"),
        (0.47, "领域服务层", "项目上下文｜区域模块｜方案版本｜资源血缘｜任务流程｜规则评价", "F4F6F9", "1F4D78"),
        (0.345, "智能能力层", "视觉分割｜多模态对齐｜知识检索｜生成探索｜参数优化｜跨区协同", "E8EEF5", "2E74B5"),
        (0.22, "平台支撑层", "统一身份｜权限策略｜资源管理｜任务编排｜通知审计｜接口适配", "F4F6F9", "1F4D78"),
        (0.095, "数据与基础设施层", "主数据｜样本库｜模块库｜材料库｜规则库｜成果库｜关系数据｜文件存储｜计算资源", "F4F6F9", "5E6975"),
    ]
    for y, title, sub, fill, edge in layers:
        fmini(draw, 0.135, y, 0.73, 0.10, title, sub, fill=fill, edge=edge)
    for y in (0.72, 0.595, 0.47, 0.345, 0.22):
        farrow(draw, (0.50, y), (0.50, y - 0.025), color="2E74B5")
    save_figure(image, "v2-1-figure-04-layered-architecture.png")

    image, draw = figure_base.setup_figure("平台服务边界、网络分区与受控部署拓扑")
    zones = [
        (0.02, "用户终端区", "设计人员\n负责人\n管理人员", "FFF0F2", ACCENT),
        (0.18, "统一接入区", "访问入口\n身份校验\n流量控制", "E8EEF5", "2E74B5"),
        (0.34, "平台应用区", "项目/资源\n任务/版本\n规则/审计", "E8EEF5", "2E74B5"),
        (0.50, "智能能力区", "识别/理解\n生成/优化\n检索/报告", "FFF0F2", ACCENT),
        (0.66, "数据资源区", "关系数据\n文件资料\n知识与索引", "F4F6F9", "1F4D78"),
        (0.82, "产品平台区", "主数据\n任务协同\n成果交换", "FFF0F2", ACCENT),
    ]
    for x, title, sub, fill, edge in zones:
        fmini(draw, x, 0.51, 0.14, 0.27, title, sub, fill=fill, edge=edge)
    for x in (0.16, 0.32, 0.48, 0.64, 0.80):
        farrow(draw, (x, 0.645), (x + 0.02, 0.645), color=ACCENT if x in {0.16, 0.48, 0.80} else "2E74B5")
    lower = [
        (0.18, "接入防护", "证书·会话·白名单"), (0.34, "平台运行", "配置·监控·告警"),
        (0.50, "计算与调度", "通用/专用计算·队列"), (0.66, "数据保护", "备份·恢复·生命周期"),
    ]
    for x, title, sub in lower:
        fmini(draw, x, 0.23, 0.14, 0.15, title, sub, fill="F4F6F9", edge="1F4D78")
        farrow(draw, (x + 0.07, 0.51), (x + 0.07, 0.38), color="1F4D78")
    fmini(draw, 0.18, 0.08, 0.62, 0.10, "受控交互原则", "终端不直连模型和私有存储；跨区访问经统一入口、权限校验、任务记录与审计", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-05-deployment-boundary.png")

    image, draw = figure_base.setup_figure("客室内装四级模块对象、属性与关系体系")
    fmini(draw, 0.05, 0.74, 0.74, 0.10, "一级：客室内装系统", "车型平台｜目标市场｜客室等级｜总体风格｜适用规范", fill="FFF0F2", edge=ACCENT)
    regions = [(0.05, "侧墙"), (0.18, "顶部"), (0.31, "地板"), (0.44, "端墙"), (0.57, "座椅"), (0.70, "门区")]
    for x, title in regions:
        fmini(draw, x, 0.57, 0.10, 0.11, f"二级：{title}", "区域对象", fill="E8EEF5")
    fmini(draw, 0.05, 0.39, 0.74, 0.12, "三级：可配置组件", "结构单元｜装饰单元｜照明单元｜信息单元｜设备接口｜检修单元", fill="F4F6F9", edge="1F4D78")
    fmini(draw, 0.05, 0.21, 0.74, 0.12, "四级：零件与表面", "几何边界｜材料色彩｜纹理工艺｜装配公差｜维护等级｜供应状态", fill="F4F6F9", edge="5E6975")
    for x in (0.42,):
        farrow(draw, (x, 0.74), (x, 0.68)); farrow(draw, (x, 0.57), (x, 0.51)); farrow(draw, (x, 0.39), (x, 0.33))
    fmini(draw, 0.82, 0.55, 0.15, 0.29, "对象主数据", "唯一编码\n名称与别名\n版本与状态\n适用范围\n来源与责任人", fill="FFF0F2", edge=ACCENT)
    fmini(draw, 0.82, 0.21, 0.15, 0.27, "关系模型", "组成关系\n空间邻接\n接口依赖\n规则约束\n方案引用", fill="E8EEF5", edge="1F4D78")
    fmini(draw, 0.18, 0.07, 0.50, 0.09, "公共属性：项目归属｜区域坐标｜参数集合｜接口集合｜规则版本｜成果血缘", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-06-module-hierarchy.png")

    image, draw = figure_base.setup_figure("多模态资料到可用分区对象的端到端识别流程")
    pipeline = [
        (0.02, "多源输入", "图像\n图纸\n三维\n文本属性"),
        (0.16, "质量门控", "清晰度\n完整性\n授权\n格式"),
        (0.30, "配准与标准化", "坐标对齐\n尺度归一\n视角校正\n语义清洗"),
        (0.44, "特征编码", "视觉\n几何\n文字\n空间关系"),
        (0.58, "多模态融合", "区域先验\n对象词典\n跨模态对齐\n证据加权"),
        (0.72, "候选推理", "类别\n边界\n实例\n置信度"),
        (0.86, "确认与发布", "规则复核\n人工修订\n模块映射\n版本入库"),
    ]
    for x, title, sub in pipeline:
        fmini(draw, x, 0.53, 0.12, 0.26, title, sub, fill="FFF0F2" if x in {0.02, 0.86} else "E8EEF5", edge=ACCENT if x in {0.02, 0.86} else "2E74B5")
    for x in (0.14, 0.28, 0.42, 0.56, 0.70, 0.84):
        farrow(draw, (x, 0.66), (x + 0.02, 0.66), color=ACCENT if x >= 0.70 else "2E74B5")
    controls = [
        (0.05, "数据规范", "采集·标注·抽样"), (0.24, "对象标准", "区域·模块·编码"),
        (0.43, "模型管理", "版本·阈值·适用域"), (0.62, "规则体系", "硬约束·冲突·例外"),
        (0.81, "评价反馈", "精度·边界·可用性"),
    ]
    for x, title, sub in controls:
        fmini(draw, x, 0.24, 0.14, 0.14, title, sub, fill="F4F6F9", edge="1F4D78")
    fmini(draw, 0.18, 0.08, 0.64, 0.10, "全过程证据", "原始资料、处理参数、候选结果、人工操作、规则版本和最终对象可追溯", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-07-multimodal-recognition.png")

    image, draw = figure_base.setup_figure("候选识别结果的人机协同分流与质量闭环")
    fmini(draw, 0.03, 0.50, 0.16, 0.24, "候选结果包", "类别·边界·实例\n置信度·证据\n模型与规则版本", fill="E8EEF5")
    fmini(draw, 0.24, 0.50, 0.16, 0.24, "自动质量评估", "阈值判断\n边界完整性\n对象一致性\n硬约束冲突", fill="F4F6F9", edge="1F4D78")
    farrow(draw, (0.19, 0.62), (0.24, 0.62))
    lanes = [
        (0.68, "A：高置信无冲突", "抽样确认", "EAF7EF", "2D7D46"),
        (0.45, "B：可修订候选", "局部边界与属性修订", "FFF7E6", "B7791F"),
        (0.22, "C：低置信或冲突", "重标、重识别或规则处置", "FFF0F2", ACCENT),
    ]
    for y, title, action, fill, edge in lanes:
        fmini(draw, 0.48, y, 0.21, 0.14, title, action, fill=fill, edge=edge)
        fmini(draw, 0.75, y, 0.20, 0.14, "专业人员确认", "责任人·意见·时间·结果", fill="F4F6F9", edge=edge)
        farrow(draw, (0.69, y + 0.07), (0.75, y + 0.07), color=edge)
    fortho(draw, [(0.40, 0.64), (0.44, 0.64), (0.44, 0.75), (0.48, 0.75)], color="2D7D46")
    farrow(draw, (0.40, 0.59), (0.48, 0.52), color="B7791F")
    fortho(draw, [(0.40, 0.55), (0.44, 0.55), (0.44, 0.29), (0.48, 0.29)], color=ACCENT)
    fmini(draw, 0.26, 0.06, 0.48, 0.10, "闭环更新", "确认结果进入数据集、阈值策略、对象词典和规则改进清单，不直接覆盖历史版本", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-08-human-review.png")

    image, draw = figure_base.setup_figure("多模态能力编排、工程约束与运行治理")
    lanes = [
        (0.72, "业务流程", ["设计意图", "区域任务", "候选方案", "评审确认"]),
        (0.55, "编排控制", ["输入校验", "流程拆解", "状态汇聚", "失败处置"]),
        (0.38, "智能能力", ["理解识别", "知识检索", "生成优化", "报告表达"]),
        (0.21, "工程约束", ["边界接口", "安全规范", "协同规则", "人工审批"]),
    ]
    xs = [0.20, 0.39, 0.58, 0.77]
    for y, label, cells in lanes:
        fmini(draw, 0.03, y, 0.13, 0.12, label, fill="FFF0F2" if label in {"业务流程", "工程约束"} else "E8EEF5", edge=ACCENT if label in {"业务流程", "工程约束"} else "1F4D78")
        for x, text_value in zip(xs, cells):
            fmini(draw, x, y, 0.15, 0.12, text_value, "输入·版本·状态·输出", fill="F4F6F9", edge="9AA7B5")
        for x in (0.35, 0.54, 0.73):
            farrow(draw, (x, y + 0.06), (x + 0.04, y + 0.06), color="2E74B5")
    for x in [0.275, 0.465, 0.655, 0.845]:
        draw.line([figure_base._point_pixels((x, 0.72)), figure_base._point_pixels((x, 0.33))], fill="#B5BFC9", width=3)
    fmini(draw, 0.20, 0.06, 0.72, 0.10, "统一任务记录", "项目、输入资源、业务参数、规则版本、能力版本、执行日志、输出成果和评审结论", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-09-capability-orchestration.png")

    image, draw = figure_base.setup_figure("模块化、参数化、规则化的快速设计引擎")
    columns = [
        (0.02, "设计输入", "区域对象\n模块库\n参考资料\n目标约束"),
        (0.19, "参数模型", "尺寸范围\n接口参数\nCMF参数\n适用条件"),
        (0.36, "组合与生成", "模块选配\n参数驱动\n生成探索\n变体管理"),
        (0.53, "规则求解", "硬约束\n条件规则\n冲突定位\n例外流程"),
        (0.70, "方案评价", "功能完整\n风格协调\n工程可行\n复用价值"),
        (0.87, "结果输出", "候选排序\n差异说明\n评审版本\n成果入库"),
    ]
    for x, title, sub in columns:
        fmini(draw, x, 0.50, 0.13, 0.29, title, sub, fill="FFF0F2" if x in {0.02, 0.87} else "E8EEF5", edge=ACCENT if x in {0.02, 0.87} else "2E74B5")
    for x in (0.15, 0.32, 0.49, 0.66, 0.83):
        farrow(draw, (x, 0.645), (x + 0.04, 0.645), color=ACCENT if x >= 0.66 else "2E74B5")
    controls = [
        (0.12, "对象与模板版本", "可复现输入"), (0.34, "规则与阈值版本", "可解释判定"),
        (0.56, "方案分支与比较", "保留差异"), (0.78, "评审结论与复用", "形成闭环"),
    ]
    for x, title, sub in controls:
        fmini(draw, x, 0.25, 0.18, 0.13, title, sub, fill="F4F6F9", edge="1F4D78")
    fmini(draw, 0.18, 0.07, 0.64, 0.11, "迭代原则", "任何自动生成结果必须通过规则校核与人工评审；修改形成新版本并保留来源关系", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-10-rapid-design-engine.png")

    image, draw = figure_base.setup_figure("客室区域依赖矩阵与协同优化逻辑")
    headers = [(0.20, "空间边界"), (0.34, "装配接口"), (0.48, "功能联动"), (0.62, "风格协调")]
    for x, title in headers:
        fmini(draw, x, 0.76, 0.12, 0.09, title, fill="FFF0F2", edge=ACCENT)
    rows = [
        (0.65, "顶部", ["净空/曲面", "侧墙收口", "照明风道", "顶墙过渡"]),
        (0.54, "侧墙", ["窗口/轮廓", "顶地接口", "设备检修", "色材连续"]),
        (0.43, "地板", ["通道/高差", "墙座连接", "防滑导向", "色彩分区"]),
        (0.32, "端墙", ["门洞/边界", "顶墙地收口", "设备布置", "视觉中心"]),
        (0.21, "门区", ["开闭包络", "墙地接口", "导向安全", "识别强化"]),
        (0.10, "座椅", ["人体空间", "地板连接", "通行维护", "材质呼应"]),
    ]
    colors = ["E8EEF5", "F4F6F9"]
    for idx, (y, region, cells) in enumerate(rows):
        fmini(draw, 0.05, y, 0.12, 0.085, region, f"优先级 {idx + 1}", fill="FFF0F2", edge=ACCENT)
        for (x, _), value in zip(headers, cells):
            fmini(draw, x, y, 0.12, 0.085, value, fill=colors[idx % 2], edge="9AA7B5")
        fmini(draw, 0.78, y, 0.17, 0.085, "冲突检查→综合评价", "记录责任对象与处置结论", fill="F4F6F9", edge="1F4D78")
    fsection_label(draw, 0.05, 0.88, "区域")
    fsection_label(draw, 0.78, 0.88, "协同处理")
    save_figure(image, "v2-1-figure-11-cross-region-optimization.png")

    image, draw = figure_base.setup_figure("平台功能域、核心模块与共享支撑能力")
    domains = [
        (0.03, "项目协同域", ["项目与成员", "需求与资料", "任务与进度", "通知与评审"]),
        (0.27, "智能设计域", ["分区识别", "模块化设计", "形态与CMF", "整体协同优化"]),
        (0.51, "成果知识域", ["方案版本", "评价对比", "成果检索", "知识沉淀复用"]),
        (0.75, "运营管理域", ["用户与权限", "规则能力管理", "运行监控", "日志与审计"]),
    ]
    for x, title, modules in domains:
        fmini(draw, x, 0.72, 0.21, 0.11, title, "业务功能域", fill="FFF0F2", edge=ACCENT)
        for i, module in enumerate(modules):
            fmini(draw, x, 0.56 - i * 0.13, 0.21, 0.10, module, fill="E8EEF5" if i % 2 == 0 else "F4F6F9", edge="2E74B5")
    fmini(draw, 0.10, 0.06, 0.80, 0.10, "共享支撑", "统一身份｜项目上下文｜资源与版本｜任务编排｜规则服务｜接口适配｜安全审计", fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-12-function-map.png")

    image, draw = figure_base.setup_figure("平台数据资源、治理流程与知识服务底座")
    sources = [
        (0.03, "基础资料", "车型·图纸·模型"), (0.215, "模块主数据", "层级·属性·接口"),
        (0.40, "材料与CMF", "色材·纹理·工艺"), (0.585, "规则与参数", "约束·模板·版本"),
        (0.77, "方案与运行", "成果·评价·日志"),
    ]
    for x, title, sub in sources:
        fmini(draw, x, 0.72, 0.16, 0.12, title, sub, fill="E8EEF5")
        farrow(draw, (x + 0.08, 0.72), (x + 0.08, 0.64))
    governance = ["采集授权", "分类编码", "质量检查", "版本血缘", "发布归档"]
    for i, title in enumerate(governance):
        x = 0.05 + i * 0.185
        fmini(draw, x, 0.51, 0.16, 0.11, title, "责任·状态·证据", fill="F4F6F9", edge="1F4D78")
        if i < 4:
            farrow(draw, (x + 0.16, 0.565), (x + 0.185, 0.565), color=ACCENT)
    products = [
        (0.08, "对象数据服务", "查询·关系·权限"), (0.31, "知识检索服务", "标签·语义·相似"),
        (0.54, "规则参数服务", "校核·推荐·追踪"), (0.77, "成果复用服务", "版本·评价·引用"),
    ]
    for x, title, sub in products:
        fmini(draw, x, 0.30, 0.16, 0.13, title, sub, fill="FFF0F2", edge=ACCENT)
    fmini(draw, 0.08, 0.10, 0.40, 0.12, "结构化数据与索引", "对象、关系、规则、任务、权限、元数据", fill="F4F6F9", edge="1F4D78")
    fmini(draw, 0.52, 0.10, 0.40, 0.12, "文件与模型资源", "图像、图纸、文档、三维、模型、成果包", fill="F4F6F9", edge="1F4D78")
    save_figure(image, "v2-1-figure-13-data-foundation.png")

    image, draw = figure_base.setup_figure("平台建设三阶段、四工作流与阶段门")
    phases = [(0.25, "第一阶段：研究与总体设计", "至 2026-09-30"), (0.49, "第二阶段：构建与检测", "至 2026-12-31"), (0.73, "第三阶段：集成与交付", "至 2027-03-31")]
    for x, title, sub in phases:
        fmini(draw, x, 0.76, 0.22, 0.10, title, sub, fill="FFF0F2", edge=ACCENT)
    workstreams = [
        (0.61, "标准与数据", ["对象标准\n数据清单", "样本与知识库\n质量治理", "数据固化\n移交与维护"]),
        (0.45, "关键技术", ["技术路线\n原型与指标", "识别/生成/规则\n迭代验证", "场景优化\n能力定版"]),
        (0.29, "平台系统", ["架构与接口\n功能原型", "功能开发\n能力接入", "产品平台联调\n运行完善"]),
        (0.13, "测试与交付", ["评审基线\n用例设计", "分层测试\n检测整改", "验收演练\n成果交付"]),
    ]
    for y, label, cells in workstreams:
        fmini(draw, 0.04, y, 0.16, 0.12, label, "并行推进", fill="E8EEF5", edge="1F4D78")
        for (x, _, _), cell in zip(phases, cells):
            fmini(draw, x, y, 0.22, 0.12, cell, fill="F4F6F9", edge="9AA7B5")
    for x, gate in [(0.25, "阶段门 G1\n研究报告与总体设计评审"), (0.49, "阶段门 G2\n系统与检测结果评审"), (0.73, "阶段门 G3\n集成成果与交付验收")]:
        fmini(draw, x, 0.02, 0.22, 0.08, gate, fill="FFF0F2", edge=ACCENT)
    save_figure(image, "v2-1-figure-14-technical-roadmap.png")

    image, draw = figure_base.setup_figure("需求、设计、验证证据与验收活动对应体系")
    headers = [(0.04, "建设定义"), (0.36, "验证证据"), (0.68, "验收活动")]
    for x, title in headers:
        fmini(draw, x, 0.77, 0.28, 0.09, title, fill="FFF0F2", edge=ACCENT)
    rows = [
        (0.64, "合同目标与交付范围", "要求追踪矩阵·成果清单", "资料完整性与范围核查"),
        (0.51, "业务场景与用户任务", "场景脚本·操作记录·评审意见", "典型项目端到端验证"),
        (0.38, "总体架构与接口边界", "部署记录·接口报告·安全记录", "系统集成与安全验证"),
        (0.25, "功能模块与业务规则", "功能用例·规则结果·缺陷闭环", "模块功能与异常验证"),
        (0.12, "算法、数据与评价指标", "数据集版本·指标报告·人工复核", "技术指标与适用域验证"),
    ]
    for y, definition, evidence, acceptance in rows:
        fmini(draw, 0.04, y, 0.28, 0.10, definition, fill="E8EEF5", edge="2E74B5")
        fmini(draw, 0.36, y, 0.28, 0.10, evidence, fill="F4F6F9", edge="1F4D78")
        fmini(draw, 0.68, y, 0.28, 0.10, acceptance, fill="FFF0F2", edge=ACCENT)
        farrow(draw, (0.32, y + 0.05), (0.36, y + 0.05))
        farrow(draw, (0.64, y + 0.05), (0.68, y + 0.05), color=ACCENT)
    draw.text(figure_base._point_pixels((0.50, 0.055)), "每一项验收结论均应能够反查需求、版本、执行记录、责任人和整改状态", font=figure_base.figure_font(23), fill="#5E6975", anchor="mm")
    save_figure(image, "v2-1-figure-15-verification-system.png")


def validate_source(source: str):
    figures = [int(x) for x in re.findall(r"!\[图\s*(\d+)\s+", source)]
    tables = [int(x) for x in re.findall(r"^表\s*(\d+)\s+", source, flags=re.MULTILINE)]
    if figures != list(range(1, len(figures) + 1)):
        raise ValueError(f"figure numbering is not sequential: {figures}")
    if tables != list(range(1, len(tables) + 1)):
        raise ValueError(f"table numbering is not sequential: {tables}")


def build():
    source = SOURCE_PATH.read_text(encoding="utf-8")
    validate_source(source)
    toc_entries = extract_toc_entries(source)
    generate_figures()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_document(doc)
    configure_header_footer(doc.sections[0])
    props = doc.core_properties
    props.title = DOC_TITLE
    props.subject = "轨道交通客运装备内装模块分区识别与快速设计平台建设指导"
    props.author = "轨道客室智能设计平台项目组"
    props.keywords = "轨道交通, 客室内装, 模块分区, 多模态识别, 快速设计, 平台架构"
    props.comments = f"文档编号 {DOC_ID}；版本 {DOC_VERSION}；建设指导评审稿"

    add_cover(doc)
    add_front_matter(doc, toc_entries)
    parse_markdown(doc, source, toc_entries)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
